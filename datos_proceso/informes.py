# datos_proceso/informes.py — Documentos de datos de proceso (fase 3): PDF del lote y Excel de uno o varios lotes.
#
# Solo da formato: los datos salen de services.datos_informe() y services.filas_lotes_excel().
# El informe sigue al diseñador: una sección por etapa (datos y gráficas por actividad · sección) y, si se piden,
# gráficas combinadas de varios parámetros (apiladas con el mismo eje, o superpuestas escaladas a 0–100 %).
#   pdf_lote(datos, emitido_por)   → bytes del PDF (carta horizontal; encabezado, pie con huella y "página X de Y")
#   (el PDF incluye «Retrasos contra lo planeado»: services.datos_informe()['retrasos'], de procesos.services.analisis_retrasos)
#   excel_lote(datos, emitido_por) → bytes del .xlsx (Resumen, Datos, Fuera de límites, Eventos GMP, Retrasos, Muestreos tarde,
#                                     Adicionales, Correcciones)
#   excel_lotes(...)               → bytes del .xlsx de varios lotes (Datos en formato largo e Indicadores)

from io import BytesIO
from xml.sax.saxutils import escape as xml_escape

from django.utils import timezone

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.graphics.shapes import Circle, Drawing, Line, PolyLine, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas as canvas_pdf
from reportlab.platypus import (CondPageBreak, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

# Colores (los mismos tonos de la pantalla de gráficas en fondo blanco)
AZUL = colors.HexColor('#2a78d6')
TINTA = colors.HexColor('#0f172a')
TENUE = colors.HexColor('#64748b')
REJILLA = colors.HexColor('#e2e8f0')
ROJO = colors.HexColor('#dc2626')
AMBAR = colors.HexColor('#d97706')
REFERENCIA = colors.HexColor('#334155')
FONDO_ENCABEZADO = colors.HexColor('#e8f0fb')
COLUMNAS_POR_TABLA = 8
# Colores de las gráficas combinadas superpuestas (uno por parámetro; mismos tonos que la pantalla)
COLORES_COMBINADAS = ['#2a78d6', '#eb6834', '#1baf7a', '#eda100', '#e87ba4', '#008300']


def _xml(texto):
    """Texto libre (escrito por usuarios) seguro dentro de un Paragraph de reportlab; respeta saltos de línea."""
    return xml_escape(texto or '').replace('\n', '<br/>')


def _local(momento, formato='%d/%m/%Y %H:%M'):
    return timezone.localtime(momento).strftime(formato) if momento else '—'


def _horas(h):
    if h is None:
        return '—'
    return '0.0' if abs(h) < 0.05 else f"{h:.1f}"   # sin "-0.0"


UMBRAL_RETRASO_H = 5 / 60   # menos de 5 minutos no se cuenta como retraso (igual que la pantalla /retrasos/)
VERDE = colors.HexColor('#16a34a')
GRIS_HEREDADO = colors.HexColor('#cbd5e1')


def _duracion(h, signo=False):
    """3.5 → '3 h 30 min'; 50 → '2 d 2 h'; con signo: '+3 h', '−45 min'. Igual que la pantalla de retrasos."""
    if h is None:
        return '—'
    total = round(abs(h) * 60)
    if total < 5:
        return '0'
    dias, horas, minutos = total // 1440, (total % 1440) // 60, total % 60
    if dias:
        texto = f"{dias} d" + (f" {horas} h" if horas else '')
    elif horas:
        texto = f"{horas} h" + (f" {minutos} min" if minutos else '')
    else:
        texto = f"{minutos} min"
    return (('+' if h > 0 else '−') + texto) if signo else texto


def _texto_espera(a, ahora):
    if a['espera_h'] > UMBRAL_RETRASO_H:
        if a['inicio_real']:
            return f"Inició {_duracion(a['espera_h'])} tarde"
        if a['estado'] == 'PENDIENTE' and a['inicio'] <= ahora:
            return f"Lleva {_duracion(a['espera_h'])} sin iniciar"
        return f"Se reprogramó {_duracion(a['espera_h'])} después"
    if a['espera_h'] < -UMBRAL_RETRASO_H:
        return f"Inició {_duracion(a['espera_h'])} antes"
    return ''


def _texto_duracion(a):
    if abs(a['duracion_extra_h']) <= UMBRAL_RETRASO_H:
        return ''
    if a['duracion_extra_h'] < 0:
        return f"Duró {_duracion(a['duracion_extra_h'])} menos"
    return f"Duró {_duracion(a['duracion_extra_h'])} más" if a['fin_real'] else f"Va {_duracion(a['duracion_extra_h'])} arriba de su duración"


def _cascada_retrasos(ruta, ancho, origen_id):
    """Cascada de la ruta: cada fila arranca en el retraso heredado y muestra lo que nació al iniciar y por duración."""
    fila, etiqueta, valor = 0.62 * cm, 6.6 * cm, 5.2 * cm
    pista = ancho - etiqueta - valor - 0.4 * cm
    alto = fila * len(ruta) + 0.3 * cm
    d = Drawing(ancho, alto)
    valores = [v for a in ruta for v in (max(0, a['heredado_h']), a['retraso_inicio_h'], a['retraso_fin_h'])]
    minimo, maximo = min(0, *valores), max(0, *valores)
    rango = (maximo - minimo) or 1
    x0 = etiqueta + 0.2 * cm
    x = lambda h: x0 + (h - minimo) / rango * pista
    for i, a in enumerate(ruta):
        y = alto - (i + 1) * fila
        if a['id'] == origen_id:
            d.add(Rect(0, y, ancho, fila, fillColor=colors.HexColor('#fef3c7'), strokeColor=None))
        d.add(Line(0, y, ancho, y, strokeColor=REJILLA, strokeWidth=0.4))
        nombre = f"[{a['paso']}] {a['nombre']}"
        d.add(String(3, y + fila * 0.55, nombre if len(nombre) <= 44 else nombre[:43] + '…', fontName='Helvetica-Bold', fontSize=7,
                     fillColor=TINTA))
        d.add(String(3, y + fila * 0.17, a['estado_texto'], fontName='Helvetica', fontSize=6, fillColor=TENUE))
        d.add(Rect(x0, y + fila * 0.2, pista, fila * 0.6, fillColor=colors.HexColor('#f1f5f9'), strokeColor=None))
        partida = max(0, a['heredado_h'])
        if partida > UMBRAL_RETRASO_H:
            d.add(Rect(x(0), y + fila * 0.42, x(partida) - x(0), fila * 0.16, fillColor=GRIS_HEREDADO, strokeColor=None))
        for desde, hasta, color in ((partida, a['retraso_inicio_h'], AMBAR), (a['retraso_inicio_h'], a['retraso_fin_h'], ROJO)):
            if abs(hasta - desde) > UMBRAL_RETRASO_H:
                izq = x(min(desde, hasta))
                d.add(Rect(izq, y + fila * 0.25, max(1.2, abs(x(hasta) - x(desde))), fila * 0.5,
                           fillColor=VERDE if hasta < desde else color, strokeColor=None))
        d.add(Line(x(0), y + fila * 0.1, x(0), y + fila * 0.9, strokeColor=TENUE, strokeWidth=0.6))
        total = a['retraso_fin_h']
        d.add(String(ancho - 3, y + fila * 0.55, 'A tiempo' if abs(total) <= UMBRAL_RETRASO_H else _duracion(total, True),
                     fontName='Helvetica-Bold', fontSize=7, textAnchor='end',
                     fillColor=ROJO if total > 24 else AMBAR if total > UMBRAL_RETRASO_H else VERDE if total < -UMBRAL_RETRASO_H else TENUE))
        propio = ' · '.join(t for t in (a['texto_espera'], a['texto_duracion']) if t) or 'Sin retraso propio'
        d.add(String(ancho - 3, y + fila * 0.17, propio, fontName='Helvetica', fontSize=6, fillColor=TENUE, textAnchor='end'))
    return d


def _numero(v, decimales=6):
    if v is None:
        return '—'
    if isinstance(v, (int, float)):
        return f"{v:.{decimales}f}".rstrip('0').rstrip('.') if decimales else f"{v:.0f}"
    return str(v)


def _meta(i):
    if i['meta_min'] is None and i['meta_max'] is None:
        return '—'
    return f"{_numero(i['meta_min']) if i['meta_min'] is not None else '−∞'} – {_numero(i['meta_max']) if i['meta_max'] is not None else '∞'}"


def _titulo_documento(datos):
    return f"Informe de datos de proceso · Lote {datos['lote'].identificador_lote}" + (f" · Etapa {datos['etapa']}" if datos['etapa'] else '')


# ==============================================================================
# PDF
# ==============================================================================

def _estilos():
    base = getSampleStyleSheet()
    return {
        'titulo': ParagraphStyle('titulo', parent=base['Title'], fontName='Helvetica-Bold', fontSize=16, leading=19,
                                 alignment=TA_LEFT, textColor=TINTA, spaceAfter=2),
        'subtitulo': ParagraphStyle('subtitulo', parent=base['Normal'], fontSize=9, textColor=TENUE, spaceAfter=8),
        'seccion': ParagraphStyle('seccion', parent=base['Heading2'], fontName='Helvetica-Bold', fontSize=11.5,
                                  textColor=AZUL, spaceBefore=10, spaceAfter=5, keepWithNext=1),
        'normal': ParagraphStyle('normal', parent=base['Normal'], fontSize=8, leading=10, textColor=TINTA),
        'celda': ParagraphStyle('celda', parent=base['Normal'], fontSize=7, leading=8.5, textColor=TINTA),
        'etapa': ParagraphStyle('etapa', parent=base['Heading1'], fontName='Helvetica-Bold', fontSize=14, leading=17, textColor=colors.white,
                                backColor=AZUL, borderPadding=(4, 6, 4, 6), spaceBefore=14, spaceAfter=6, keepWithNext=1),
        'subseccion': ParagraphStyle('subseccion', parent=base['Heading3'], fontName='Helvetica-Bold', fontSize=10, textColor=AZUL,
                                     spaceBefore=8, spaceAfter=4, keepWithNext=1),
        'bloque': ParagraphStyle('bloque', parent=base['Normal'], fontName='Helvetica-Bold', fontSize=8.5, textColor=REFERENCIA,
                                 spaceBefore=4, spaceAfter=2, keepWithNext=1),
        'cabecera': ParagraphStyle('cabecera', parent=base['Normal'], fontName='Helvetica-Bold', fontSize=7, leading=8.5, textColor=TINTA),
        'nota': ParagraphStyle('nota', parent=base['Normal'], fontSize=7, leading=9, textColor=TENUE),
    }


def _estilo_tabla(encabezados=1, zebra=True):
    comandos = [
        ('FONTNAME', (0, 0), (-1, encabezados - 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('BACKGROUND', (0, 0), (-1, encabezados - 1), FONDO_ENCABEZADO),
        ('TEXTCOLOR', (0, 0), (-1, -1), TINTA),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LINEBELOW', (0, encabezados - 1), (-1, encabezados - 1), 0.8, AZUL),
        ('LINEBELOW', (0, encabezados), (-1, -1), 0.25, REJILLA),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
    ]
    if zebra:
        comandos.append(('ROWBACKGROUNDS', (0, encabezados), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]))
    return TableStyle(comandos)


class _LienzoNumerado(canvas_pdf.Canvas):
    """Lienzo que conoce el total de páginas para escribir "Página X de Y" en el pie."""

    def __init__(self, *args, encabezado=None, **kwargs):
        super().__init__(*args, **kwargs)
        self._paginas = []
        self._encabezado = encabezado

    def showPage(self):
        self._paginas.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._paginas)
        for estado in self._paginas:
            self.__dict__.update(estado)
            self._encabezado(self, total)
            super().showPage()
        super().save()


def _grafica(serie, ancho, alto, x_rango=None, titulo=None, curva_si=True):
    """Gráfica de un parámetro del lote contra la edad, con curva esperada, valor esperado y límites.
    `x_rango` fija el eje X (gráficas combinadas apiladas, todas alineadas)."""
    c = serie['columna']
    ref = c.get('referencia') or {}
    lim = c.get('limites') or {}
    puntos = serie['puntos']
    curva = [(float(x), float(y)) for x, y in (ref.get('curva') or [])] if curva_si else []
    objetivo = float(ref['objetivo']) if ref.get('objetivo') not in (None, '') else None
    lineas = [(float(lim[k]), etiqueta, color) for k, etiqueta, color in (
        ('especificacion_max', 'Esp. máx', ROJO), ('especificacion_min', 'Esp. mín', ROJO),
        ('alerta_max', 'Alerta máx', AMBAR), ('alerta_min', 'Alerta mín', AMBAR)) if lim.get(k) not in (None, '')]

    xs = [p[0] for p in puntos]
    if x_rango:
        x_min, x_max = x_rango
    else:
        x_min, x_max = min(0, min(xs)), max(xs)
        x_max = max(x_max + max(2, (x_max - x_min) * 0.06), 24)
    curva_visible = [(x, y) for x, y in curva if x_min <= x <= x_max]
    ys = [p[1] for p in puntos] + [y for y, _, _ in lineas] + ([objetivo] if objetivo is not None else []) + [y for _, y in curva_visible]
    y_min, y_max = min(ys), max(ys)
    holgura = (y_max - y_min) * 0.1 or abs(y_max) * 0.1 or 1
    y_min, y_max = y_min - holgura, y_max + holgura

    d = Drawing(ancho, alto)
    m_izq, m_der, m_arr, m_aba = 34, 8, 18, 24
    au, al = ancho - m_izq - m_der, alto - m_arr - m_aba
    X = lambda v: m_izq + (v - x_min) / (x_max - x_min) * au
    Y = lambda v: m_aba + (v - y_min) / (y_max - y_min) * al
    d.add(String(m_izq, alto - 11, titulo or (f"{c['nombre']}" + (f" ({c['unidad']})" if c['unidad'] else '')),
                 fontName='Helvetica-Bold', fontSize=8.5, fillColor=TINTA))
    d.add(Rect(m_izq, m_aba, au, al, strokeColor=REJILLA, fillColor=colors.white, strokeWidth=0.6))

    def marcas(a, b, n):
        import math
        paso0 = (b - a) / n or 1
        base = 10 ** math.floor(math.log10(paso0))
        paso = next((k * base for k in (1, 2, 2.5, 5, 10) if k * base >= paso0), paso0)
        v = math.ceil(a / paso) * paso
        while v <= b + paso * 1e-9:
            yield round(v, 10)
            v += paso

    for v in marcas(y_min, y_max, 4):
        d.add(Line(m_izq, Y(v), m_izq + au, Y(v), strokeColor=REJILLA, strokeWidth=0.4))
        d.add(String(m_izq - 3, Y(v) - 2.5, _numero(v, 3), fontName='Helvetica', fontSize=6, fillColor=TENUE, textAnchor='end'))
    for v in marcas(x_min, x_max, 6):
        d.add(String(X(v), m_aba - 9, _numero(v, 1), fontName='Helvetica', fontSize=6, fillColor=TENUE, textAnchor='middle'))
    d.add(String(m_izq + au / 2, 2, f"Edad de la etapa {serie.get('etapa', '')} (h)".replace('  ', ' ') if serie.get('eje') == 'ETAPA' else 'Edad del lote (h)',
                 fontName='Helvetica', fontSize=6.5, fillColor=TENUE, textAnchor='middle'))

    for y, etiqueta, color in lineas:
        d.add(Line(m_izq, Y(y), m_izq + au, Y(y), strokeColor=color, strokeWidth=0.8, strokeDashArray=[3, 2]))
        d.add(String(m_izq + au - 2, Y(y) + 2, f"{etiqueta} {_numero(y, 3)}", fontName='Helvetica', fontSize=5.5, fillColor=color, textAnchor='end'))
    if objetivo is not None:
        d.add(Line(m_izq, Y(objetivo), m_izq + au, Y(objetivo), strokeColor=REFERENCIA, strokeWidth=0.7, strokeDashArray=[1, 2]))
    if len(curva_visible) > 1:
        d.add(PolyLine([v for x, y in curva_visible for v in (X(x), Y(y))], strokeColor=REFERENCIA, strokeWidth=1,
                       strokeDashArray=[4, 2]))
    if len(puntos) > 1:
        d.add(PolyLine([v for x, y, _ in puntos for v in (X(x), Y(y))], strokeColor=AZUL, strokeWidth=1.4))
    for x, y, fuera in puntos:
        d.add(Circle(X(x), Y(y), 2.2, fillColor=ROJO if fuera else AZUL, strokeColor=colors.white, strokeWidth=0.6))
    return d


def _grafica_superpuesta(combinadas, ancho, alto):
    """Varios parámetros en una gráfica: cada uno escalado de su mínimo (0 %) a su máximo (100 %) contra la edad del lote."""
    d = Drawing(ancho, alto)
    m_izq, m_der, m_arr, m_aba = 34, 10, 30, 24
    au, al = ancho - m_izq - m_der, alto - m_arr - m_aba
    xs = [p[0] for c in combinadas for p in c['puntos']]
    x_min, x_max = min(0, min(xs)), max(xs)
    x_max = max(x_max + max(2, (x_max - x_min) * 0.04), 24)
    X = lambda v: m_izq + (v - x_min) / (x_max - x_min) * au
    Y = lambda v: m_aba + v / 100 * al
    d.add(Rect(m_izq, m_aba, au, al, strokeColor=REJILLA, fillColor=colors.white, strokeWidth=0.6))
    for v in (0, 25, 50, 75, 100):
        d.add(Line(m_izq, Y(v), m_izq + au, Y(v), strokeColor=REJILLA, strokeWidth=0.4))
        d.add(String(m_izq - 3, Y(v) - 2.5, f"{v} %", fontName='Helvetica', fontSize=6, fillColor=TENUE, textAnchor='end'))
    import math
    paso0 = (x_max - x_min) / 8
    base = 10 ** math.floor(math.log10(paso0))
    paso = next((k * base for k in (1, 2, 2.5, 5, 10) if k * base >= paso0), paso0)
    v = math.ceil(x_min / paso) * paso
    while v <= x_max:
        d.add(String(X(v), m_aba - 9, _numero(v, 1), fontName='Helvetica', fontSize=6, fillColor=TENUE, textAnchor='middle'))
        v += paso
    d.add(String(m_izq + au / 2, 2, 'Edad del lote (h)', fontName='Helvetica', fontSize=6.5, fillColor=TENUE, textAnchor='middle'))
    x_leyenda = m_izq
    for k, c in enumerate(combinadas):
        color = colors.HexColor(COLORES_COMBINADAS[k % len(COLORES_COMBINADAS)])
        valores = [p[1] for p in c['puntos']]
        v_min, v_max = min(valores), max(valores)
        escala = (lambda val, a=v_min, b=v_max: (val - a) / (b - a) * 100 if b > a else 50)
        pts = [(X(x), Y(escala(y))) for x, y, _ in c['puntos']]
        if len(pts) > 1:
            d.add(PolyLine([v for p in pts for v in p], strokeColor=color, strokeWidth=1.4))
        for (px, py), (_, _, fuera) in zip(pts, c['puntos']):
            d.add(Circle(px, py, 2, fillColor=ROJO if fuera else color, strokeColor=colors.white, strokeWidth=0.5))
        col = c['columna']
        texto = f"{col['nombre']} ({c['etapa']}): {_numero(v_min, 3)} → {_numero(v_max, 3)} {col['unidad']}".strip()
        d.add(Line(x_leyenda, alto - 12, x_leyenda + 12, alto - 12, strokeColor=color, strokeWidth=2))
        d.add(String(x_leyenda + 15, alto - 14.5, texto, fontName='Helvetica', fontSize=6.5, fillColor=TINTA))
        x_leyenda += 22 + len(texto) * 3.3
    return d


def _tabla_datos(h, e, etapa, ancho_util):
    """Tabla de las tomas de una etapa (las columnas de parámetros se reparten en bloques para que quepan)."""
    columnas, tomas = etapa['columnas'], etapa['tomas']
    if not tomas:
        h.append(Paragraph('No hay datos capturados en esta etapa.', e['normal']))
        return
    for inicio in range(0, max(len(columnas), 1), COLUMNAS_POR_TABLA):
        bloque = columnas[inicio:inicio + COLUMNAS_POR_TABLA]
        if not bloque:
            break
        cabecera = [Paragraph(t, e['cabecera']) for t in ['Tarea', 'Toma', 'Edad lote (h)', 'Edad etapa (h)']] + \
                   [Paragraph(f"<b>{c['nombre']}</b>" + (f"<br/><font color='#64748b'>{c['unidad']}</font>" if c['unidad'] else '')
                              + (" <font color='#6250d6'>ƒ</font>" if c['calculado'] else ''), e['celda']) for c in bloque]
        filas, estilos_extra, fijas = [cabecera], [], 4
        for k, tm in enumerate(tomas, start=1):
            fila = [Paragraph(tm['tarea'] + (' <font color="#d97706">(adicional)</font>' if tm['adicional'] else ''), e['celda']),
                    _local(tm['momento']), _horas(tm['edad_lote_h']), _horas(tm['edad_etapa_h'])]
            for j, c in enumerate(bloque):
                v = tm['valores'].get(c['parametro'])
                texto = '' if not v else (v['texto'] or '') + ('*' if v['fuera_especificacion'] else '') + ('†' if v['corregido'] else '')
                fila.append(texto)
                if v and (v['fuera_especificacion'] or v['fuera_alerta']):
                    color = ROJO if v['fuera_especificacion'] else AMBAR
                    estilos_extra += [('TEXTCOLOR', (fijas + j, k), (fijas + j, k), color),
                                      ('FONTNAME', (fijas + j, k), (fijas + j, k), 'Helvetica-Bold')]
            filas.append(fila)
        ancho_fijo = 6.2 * cm + 2.6 * cm + 1.6 * cm + 1.6 * cm
        ancho_col = (ancho_util - ancho_fijo) / len(bloque)
        t = Table(filas, colWidths=[6.2 * cm, 2.6 * cm, 1.6 * cm, 1.6 * cm] + [ancho_col] * len(bloque), repeatRows=1)
        t.setStyle(_estilo_tabla())
        t.setStyle(TableStyle([('ALIGN', (fijas - 2, 1), (-1, -1), 'RIGHT')] + estilos_extra))
        h.append(t)
        h.append(Spacer(1, 6))


def _rejilla_graficas(series, ancho_util):
    """Gráficas de dos en dos por renglón."""
    ancho_g, alto_g = ancho_util / 2 - 0.3 * cm, 6.1 * cm
    renglones = []
    for k in range(0, len(series), 2):
        par = series[k:k + 2]
        renglones.append([_grafica(s, ancho_g, alto_g) for s in par] + ([''] if len(par) == 1 else []))
    t = Table(renglones, colWidths=[ancho_util / 2] * 2)
    t.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOTTOMPADDING', (0, 0), (-1, -1), 4)]))
    return t


def _seccion_retrasos(h, e, r, ancho_util, por_etapa):
    """Retrasos contra lo planeado: resumen, cascada de la ruta, actividades con retraso propio y muestreos fuera de tolerancia."""
    ahora = r['ahora']
    estados = {'PENDIENTE': 'Pendiente', 'EN_PROGRESO': 'En curso', 'COMPLETADA': 'Completada', 'OMITIDA': 'Omitida'}
    for a in r['todas'] + r['muestreos']:
        a['estado_texto'] = estados.get(a['estado'], a['estado']) + (f" · {a['etapa']}" if a.get('etapa') else '')
        if 'espera_h' in a:
            a['texto_espera'], a['texto_duracion'] = _texto_espera(a, ahora), _texto_duracion(a)

    h.append(CondPageBreak(7 * cm))
    h.append(Paragraph('Retrasos contra lo planeado', e['seccion']))
    terminado = r['lote']['terminado']
    total = r['retraso_horas']
    resumen = (f"<b>{'Retraso al terminar' if terminado else 'Retraso estimado al terminar'}: "
               f"{'a tiempo' if abs(total) <= UMBRAL_RETRASO_H else _duracion(total, True)}</b> · fin planeado {_local(r['fin_base'])} → "
               f"{'real' if terminado else 'estimado'} {_local(r['fin_estimado'])}.")
    if r['origen']:
        o = r['origen']
        a = next((x for x in r['todas'] if x['id'] == o['fase']), None)
        causa = (a['texto_espera'] if o['tipo'] == 'INICIO' else a['texto_duracion']) if a else ''
        resumen += f" Origen principal: <b>[{_xml(o['paso'])}] {_xml(o['nombre'])}</b> ({causa.lower()})."
    if not terminado:
        resumen += (f" Lote abierto: estimado al {_local(ahora)} (lo que ya debió iniciar se cuenta desde ese momento); "
                    'por eso no forma parte de la huella.')
    h.append(Paragraph(resumen, e['normal']))

    por_id = {a['id']: a for a in r['todas']}
    ruta = [por_id[i] for i in r['ruta'] if i in por_id]
    if ruta:
        h.append(Paragraph('Cómo se acumuló el retraso (ruta que fija el fin del lote)', e['subseccion']))
        h.append(_cascada_retrasos(ruta, ancho_util, r['origen']['fase'] if r['origen'] else None))
        h.append(Paragraph('Cada fila arranca en el retraso heredado de la actividad anterior (gris) y muestra lo que nació en ella: '
                           'ámbar al iniciar (inició tarde, sigue sin iniciar o se reprogramó), rojo por durar más de lo planeado, '
                           'verde tiempo recuperado. A la derecha, el retraso con que termina.', e['nota']))

    def motivos(a):
        partes = []
        if a.get('reprogramacion'):
            partes.append(f"<i>Reprogramada:</i> {_xml(a['reprogramacion']['motivo'])}"
                          + (f" — {_xml(a['reprogramacion']['notas'])}" if a['reprogramacion']['notas'] else ''))
        for m in a['motivos']:
            partes.append(f"<i>{_xml(m['que'])} ({_local(m['momento'])}, {_xml(m['actor'])}):</i> {_xml(m['texto'])}")
        for ev in a['eventos']:
            partes.append(f"<b>{_xml(ev['folio'])}</b> {_xml(ev['tipo'])}: {_xml(ev['titulo'])}"
                          + (' <font color="#d97706">(abierto)</font>' if ev['estado'] == 'ABIERTO' else ''))
        return Paragraph('<br/>'.join(partes) or '—', e['celda'])

    def horas(v):
        return '—' if abs(v) <= UMBRAL_RETRASO_H else _duracion(v, True)

    con_retraso = [a for a in r['actividades'] if a['espera_h'] > UMBRAL_RETRASO_H or a['duracion_extra_h'] > UMBRAL_RETRASO_H
                   or a['motivos'] or a['eventos'] or a.get('reprogramacion')]
    h.append(Paragraph('Actividades que agregaron retraso' + (' (de esta etapa)' if por_etapa else ''), e['subseccion']))
    if con_retraso:
        filas = [['Actividad', 'Inicio plan / real', 'Fin plan / real', 'Heredado', 'Nació al iniciar', 'Nació por duración',
                  'Al terminar', 'Motivos y eventos']]
        for a in con_retraso:
            real = lambda plan, valor, es_real: Paragraph(f"{_local(plan)}<br/><font color='#64748b'>{'Real' if es_real else 'Estimado'}: "
                                                          f"{_local(valor)}</font>", e['celda'])
            filas.append([Paragraph(f"<b>[{_xml(a['paso'])}]</b> {_xml(a['nombre'])}<br/><font color='#64748b'>{_xml(a['estado_texto'])}"
                                    f"{' · en la ruta' if a['en_ruta'] else ''}</font>", e['celda']),
                          real(a['base_inicio'], a['inicio'], a['inicio_real']), real(a['base_fin'], a['fin'], a['fin_real']),
                          horas(a['heredado_h']), horas(a['espera_h']), horas(a['duracion_extra_h']), horas(a['retraso_fin_h']), motivos(a)])
        t = Table(filas, colWidths=[4.8 * cm, 2.9 * cm, 2.9 * cm, 1.9 * cm, 2.1 * cm, 2.3 * cm, 1.9 * cm, ancho_util - 18.8 * cm], repeatRows=1)
        t.setStyle(_estilo_tabla())
        t.setStyle(TableStyle([('VALIGN', (0, 1), (-1, -1), 'TOP'), ('ALIGN', (3, 0), (6, -1), 'RIGHT')]
                              + [('BACKGROUND', (0, k), (-1, k), colors.HexColor('#fef3c7'))
                                 for k, a in enumerate(con_retraso, start=1) if r['origen'] and a['id'] == r['origen']['fase']]))
        h.append(t)
    else:
        h.append(Paragraph('Ninguna actividad agregó retraso.', e['normal']))

    fuera = [m for m in r['muestreos'] if m['fuera_tolerancia']]
    h.append(Paragraph('Muestreos fuera de tolerancia', e['subseccion']))
    if fuera:
        filas = [['Muestreo', 'Le tocaba', 'Se tomó', 'Retraso', 'Tolerancia', 'Motivos y eventos']]
        for m in fuera:
            filas.append([Paragraph(f"<b>[{_xml(m['paso'])}]</b> {_xml(m['nombre'])}", e['celda']), _local(m['esperado']),
                          _local(m['inicio']) if m['inicio_real'] else 'Aún no se toma', _duracion(m['retraso_h'], True),
                          _duracion(m['tolerancia_h']) if m['tolerancia_h'] >= 5 / 60 else '0 h', motivos(m)])
        t = Table(filas, colWidths=[6 * cm, 3 * cm, 3 * cm, 2.4 * cm, 2.2 * cm, ancho_util - 16.6 * cm], repeatRows=1)
        t.setStyle(_estilo_tabla())
        t.setStyle(TableStyle([('VALIGN', (0, 1), (-1, -1), 'TOP')]))
        h.append(t)
    else:
        h.append(Paragraph(f"Ningún muestreo fuera de tolerancia ({len(r['muestreos'])} tomado{'s' if len(r['muestreos']) != 1 else ''} "
                           'tarde dentro de ella).' if r['muestreos'] else 'Ningún muestreo fuera de tolerancia.', e['normal']))


def pdf_lote(datos, emitido_por):
    """PDF del informe de datos de proceso de un lote (o de una etapa)."""
    e = _estilos()
    salida = BytesIO()
    lote = datos['lote']
    emitido = _local(datos['generado'], '%d/%m/%Y %H:%M:%S')

    def encabezado_y_pie(lienzo, total):
        ancho, alto = landscape(letter)
        lienzo.saveState()
        lienzo.setFillColor(AZUL)
        lienzo.rect(0, alto - 0.45 * cm, ancho, 0.45 * cm, stroke=0, fill=1)
        lienzo.setFont('Helvetica-Bold', 8)
        lienzo.setFillColor(TINTA)
        lienzo.drawString(1.2 * cm, alto - 1.05 * cm, 'MES Bioprocesos · Datos de proceso')
        lienzo.setFont('Helvetica', 8)
        lienzo.drawRightString(ancho - 1.2 * cm, alto - 1.05 * cm, _titulo_documento(datos))
        lienzo.setStrokeColor(REJILLA)
        lienzo.line(1.2 * cm, 1.25 * cm, ancho - 1.2 * cm, 1.25 * cm)
        lienzo.setFont('Helvetica', 6.5)
        lienzo.setFillColor(TENUE)
        lienzo.drawString(1.2 * cm, 0.8 * cm, f"Emitido por {emitido_por} el {emitido} · Huella SHA-256 de los datos: {datos['huella'][:16]}…")
        lienzo.drawRightString(ancho - 1.2 * cm, 0.8 * cm, f"Página {lienzo.getPageNumber()} de {total}")
        lienzo.restoreState()

    doc = SimpleDocTemplate(salida, pagesize=landscape(letter), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.6 * cm, title=_titulo_documento(datos),
                            author=emitido_por, subject='Informe de datos de proceso')
    ancho_util = doc.width
    h = []

    # Portada: datos generales
    h.append(Paragraph(_titulo_documento(datos), e['titulo']))
    h.append(Paragraph(f"Receta {datos['receta']} · emitido el {emitido} por {emitido_por}", e['subtitulo']))
    cero_lote, cero_etapa = datos['edad_cero_lote'], datos['edad_cero_etapa']
    generales = [
        ['Lote', lote.identificador_lote, 'Estado', lote.get_estado_display()],
        ['Receta', datos['receta'], 'Alcance', f"Etapa {datos['etapa']}" if datos['etapa'] else 'Todo el lote'],
        ['Inicio real', _local(datos['inicio_real']), 'Fin real', _local(datos['fin_real'])],
        ['Edad 0 del lote', f"{_local(cero_lote['momento']) if cero_lote and cero_lote['momento'] else '—'}"
                            f"{' (estimada)' if cero_lote and cero_lote['estimada'] else ''}",
         'Origen', Paragraph(cero_lote['origen'] if cero_lote else '—', e['celda'])],
    ]
    if datos['etapa']:
        generales.append(['Edad 0 de la etapa', f"{_local(cero_etapa['momento']) if cero_etapa and cero_etapa['momento'] else '—'}"
                                                f"{' (estimada)' if cero_etapa and cero_etapa['estimada'] else ''}",
                          'Origen', Paragraph(cero_etapa['origen'] if cero_etapa else '—', e['celda'])])
    generales.append(['Tomas de datos', str(len(datos['tomas'])), 'Fuera de límites / correcciones',
                      f"{len(datos['fuera'])} / {len(datos['correcciones'])}"])
    abiertos = sum(1 for ev in datos['eventos'] if ev['estado'] == 'Abierto')
    generales.append(['Eventos GMP', f"{len(datos['eventos'])}" + (f" ({abiertos} abierto{'s' if abiertos != 1 else ''})" if abiertos else ''),
                      'Desviaciones', ', '.join(sorted({ev['numero_desviacion'] for ev in datos['eventos'] if ev['numero_desviacion']})) or '—'])
    r = datos.get('retrasos')
    if r:
        origen = r['origen']
        generales.append(['Retraso vs. planeado',
                          ('A tiempo' if abs(r['retraso_horas']) <= UMBRAL_RETRASO_H else _duracion(r['retraso_horas'], True))
                          + ('' if r['lote']['terminado'] else ' (estimado)'),
                          'Origen del retraso',
                          Paragraph(_xml(f"[{origen['paso']}] {origen['nombre']} · "
                                         f"{'inicio tarde' if origen['tipo'] == 'INICIO' else 'duró más'} {_duracion(origen['horas'])}")
                                    if origen else '—', e['celda'])])
    t = Table(generales, colWidths=[3.6 * cm, ancho_util / 2 - 3.6 * cm, 4.4 * cm, ancho_util / 2 - 4.4 * cm])
    t.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'), ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8), ('TEXTCOLOR', (0, 0), (0, -1), TENUE), ('TEXTCOLOR', (2, 0), (2, -1), TENUE),
        ('LINEBELOW', (0, 0), (-1, -1), 0.25, REJILLA), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOX', (0, 0), (-1, -1), 0.6, REJILLA),
    ]))
    h.append(t)

    # Indicadores
    if datos['indicadores']:
        h.append(Paragraph('Indicadores', e['seccion']))
        filas = [['Indicador', 'Etapa', 'Tipo', 'Valor', 'Unidad', 'Meta', 'Resultado', 'Nota']]
        estilos_extra = []
        for k, i in enumerate(datos['indicadores'], start=1):
            filas.append([i['nombre'], i['etapa'], i['tipo'], _numero(i['valor']), i['unidad'], _meta(i), i['estado'] or '—',
                          Paragraph(i['nota'] or '', e['celda'])])
            if i['estado'] in ('Bajo meta', 'Sobre meta'):
                estilos_extra.append(('TEXTCOLOR', (6, k), (6, k), ROJO))
        t = Table(filas, colWidths=[4.6 * cm, 2.8 * cm, 4.4 * cm, 2.2 * cm, 2.2 * cm, 2.8 * cm, 2.2 * cm, ancho_util - 21.2 * cm], repeatRows=1)
        t.setStyle(_estilo_tabla())
        for cmd in estilos_extra:
            t.setStyle(TableStyle([cmd, ('FONTNAME', cmd[1], cmd[2], 'Helvetica-Bold')]))
        h.append(t)

    # Una sección por etapa, como en el diseñador: sus datos y sus gráficas por actividad · sección
    if not datos['etapas']:
        h.append(Paragraph('Datos capturados', e['seccion']))
        h.append(Paragraph('No hay datos capturados.', e['normal']))
    for etapa in datos['etapas']:
        h.append(CondPageBreak(6 * cm))
        cero = etapa['edad_cero']
        h.append(Paragraph(f"Etapa {etapa['etapa']}", e['etapa']))
        h.append(Paragraph(f"{len(etapa['tomas'])} toma{'s' if len(etapa['tomas']) != 1 else ''} · edad 0 de la etapa: "
                           f"{_local(cero['momento']) if cero and cero['momento'] else '—'}"
                           f"{' (estimada)' if cero and cero['estimada'] else ''}{' · ' + cero['origen'] if cero else ''}", e['nota']))
        h.append(Paragraph('Datos capturados', e['subseccion']))
        _tabla_datos(h, e, etapa, ancho_util)
        if etapa['bloques']:
            h.append(CondPageBreak(8 * cm))
            h.append(Paragraph('Gráficas', e['subseccion']))
            for bloque in etapa['bloques']:
                h.append(KeepTogether([Paragraph(bloque['titulo'], e['bloque']), _rejilla_graficas(bloque['series'], ancho_util)]))
    if datos['tomas']:
        h.append(Paragraph('* Fuera de especificación (rojo) · en ámbar, fuera del rango de alerta · † valor corregido (ver «Correcciones») · '
                           'ƒ calculado con su fórmula. Gráficas: línea azul = datos del lote (punto rojo: fuera de especificación) · '
                           'gris discontinua = curva esperada · gris punteada = valor esperado · rojo y ámbar = límites.', e['nota']))

    # Gráficas combinadas (varios parámetros contra la edad del lote)
    if datos['combinadas']:
        h.append(CondPageBreak(9 * cm))
        apiladas = datos['vista_combinadas'] == 'apiladas'
        h.append(Paragraph('Gráficas combinadas' + (' (apiladas, mismo eje)' if apiladas else ' (superpuestas, 0–100 %)'), e['seccion']))
        if apiladas:
            xs = [p[0] for c in datos['combinadas'] for p in c['puntos']]
            x_min, x_max = min(0, min(xs)), max(xs)
            rango = (x_min, max(x_max + max(2, (x_max - x_min) * 0.04), 24))
            n = len(datos['combinadas'])
            alto_c = max(3.2 * cm, min(5.6 * cm, 15.5 * cm / n))   # todas caben juntas en una hoja
            graficas = []
            for c in datos['combinadas']:
                col = c['columna']
                titulo = f"{col['nombre']}" + (f" ({col['unidad']})" if col['unidad'] else '') + f" · Etapa {c['etapa']}"
                curva_lote = (col['referencia'] or {}).get('eje_edad', 'LOTE') == 'LOTE'
                graficas.append(_grafica({'columna': col, 'eje': 'LOTE', 'puntos': c['puntos']}, ancho_util, alto_c, x_rango=rango,
                                         titulo=titulo, curva_si=curva_lote))
            h.append(KeepTogether(graficas + [Paragraph('Todas con el mismo eje de edad del lote: lo que pasa a la misma edad '
                                                        'queda en la misma vertical.', e['nota'])]))
        else:
            h.append(_grafica_superpuesta(datos['combinadas'], ancho_util, 10 * cm))
            h.append(Paragraph('Cada parámetro va escalado de su valor más bajo (0 %) a su valor más alto (100 %) en este lote; '
                               'la leyenda indica su rango real. No hay dos ejes Y: compara tendencias, no magnitudes.', e['nota']))

    # Fuera de límites
    h.append(Paragraph('Valores fuera de límites', e['seccion']))
    if datos['fuera']:
        filas = [['Etapa', 'Parámetro', 'Valor', 'Límite', 'Rango permitido', 'Tarea', 'Toma', 'Edad lote (h)']]
        for f in datos['fuera']:
            filas.append([f['etapa'], f['parametro'], f"{f['valor']} {f['unidad']}".strip(), f['tipo'], f['rango'],
                          Paragraph(f['tarea'], e['celda']), _local(f['momento']), _horas(f['edad_lote_h'])])
        t = Table(filas, colWidths=[2.6 * cm, 4.2 * cm, 2.4 * cm, 2.4 * cm, 3.2 * cm, ancho_util - 20.2 * cm, 3 * cm, 2.4 * cm], repeatRows=1)
        t.setStyle(_estilo_tabla())
        h.append(t)
    else:
        h.append(Paragraph('Ningún valor quedó fuera de especificación ni del rango de alerta.', e['normal']))

    # Eventos GMP (fallas, fuera de especificación, desviaciones) con su evaluación firmada
    h.append(Paragraph('Eventos GMP', e['seccion']))
    if datos['eventos']:
        filas = [['Folio', 'Ocurrió', 'Tipo / tarea', 'Qué pasó y acción inmediata', 'Estado y evaluación', 'Firma de Calidad']]
        for ev in datos['eventos']:
            x = {k: _xml(v) if isinstance(v, str) else v for k, v in ev.items()}   # texto libre: se escapa para reportlab
            que = f"<b>{x['titulo']}</b><br/>{x['descripcion']}" + (f"<br/><i>Acción inmediata:</i> {x['accion_inmediata']}" if x['accion_inmediata'] else '')
            que += f"<br/><font color='#64748b'>Reportó: {x['reporto']} ({x['origen'].lower()})</font>"
            estado = f"<b>{x['estado']}</b>" + (f" · {x['clasificacion']}" if x['clasificacion'] else '')
            estado += (f"<br/>Desviación: <b>{x['numero_desviacion']}</b>" if x['numero_desviacion'] else '') + (f"<br/>{x['evaluacion']}" if x['evaluacion'] else '')
            firma = ev['firma']
            texto_firma = (f"<b>{_xml(firma['nombre'])}</b><br/>{_xml(firma['significado'])}<br/>{_local(firma['fecha'], '%d/%m/%Y %H:%M:%S')} · {firma['metodo']}"
                           if firma else '<font color="#dc2626">Pendiente de evaluación</font>')
            tipo = x['tipo'] + (f"<br/>{x['tarea']}" if x['tarea'] else '') + (f"<br/>Equipo {x['equipo']}" if x['equipo'] else '')
            filas.append([ev['folio'], _local(ev['ocurrio_en']), Paragraph(tipo, e['celda']), Paragraph(que, e['celda']),
                          Paragraph(estado, e['celda']), Paragraph(texto_firma, e['celda'])])
        t = Table(filas, colWidths=[2.3 * cm, 2.4 * cm, 3.6 * cm, ancho_util - 17.7 * cm, 5 * cm, 4.4 * cm], repeatRows=1)
        t.setStyle(_estilo_tabla())
        h.append(t)
    else:
        h.append(Paragraph('Sin eventos GMP registrados.', e['normal']))

    if datos.get('retrasos'):
        _seccion_retrasos(h, e, datos['retrasos'], ancho_util, bool(datos['etapa']))

    # Muestras adicionales
    if datos['adicionales']:
        h.append(Paragraph('Muestras adicionales', e['seccion']))
        filas = [['Tarea', 'Toma', 'Código', 'Motivo', 'Registró']]
        for a in datos['adicionales']:
            filas.append([Paragraph(a['tarea'], e['celda']), _local(a['momento']), a['muestras'] or '—',
                          Paragraph(a['motivo'] or '', e['celda']), a['registrado_por']])
        t = Table(filas, colWidths=[6 * cm, 2.8 * cm, 3 * cm, ancho_util - 16.8 * cm, 5 * cm], repeatRows=1)
        t.setStyle(_estilo_tabla())
        h.append(t)

    # Correcciones (bitácora) con firma electrónica
    h.append(Paragraph('Correcciones (bitácora de auditoría)', e['seccion']))
    if datos['correcciones']:
        filas = [['Fecha', 'Tarea', 'Cambio (antes → después)', 'Motivo', 'Firma electrónica']]
        for c in datos['correcciones']:
            cambios = '<br/>'.join(f"<b>{x['codigo']}</b>: {x['antes'] if x['antes'] is not None else 'sin dato'} → "
                                   f"{x['despues'] if x['despues'] is not None else 'sin dato'}" for x in c['cambios'])
            firma = c['firma']
            texto_firma = (f"<b>{firma['nombre']}</b><br/>{firma['significado']}<br/>{_local(firma['fecha'], '%d/%m/%Y %H:%M:%S')} · "
                           f"{firma['metodo']}") if firma else f"{c['actor']} (sin firma)"
            filas.append([_local(c['fecha']), Paragraph(c['tarea'], e['celda']), Paragraph(cambios, e['celda']),
                          Paragraph(c['motivo'] or '', e['celda']), Paragraph(texto_firma, e['celda'])])
        t = Table(filas, colWidths=[2.6 * cm, 5 * cm, 5 * cm, ancho_util - 19.6 * cm, 7 * cm], repeatRows=1)
        t.setStyle(_estilo_tabla())
        h.append(t)
    else:
        h.append(Paragraph('Sin correcciones.', e['normal']))

    # Firmas en papel y huella completa
    hueco = 1.5 * cm
    caja = (ancho_util - 2 * hueco) / 3
    firmas = Table([['Elaboró', '', 'Revisó', '', 'Aprobó'], [''] * 5, ['Nombre, firma y fecha', '', 'Nombre, firma y fecha', '',
                                                                        'Nombre, firma y fecha']],
                   colWidths=[caja, hueco, caja, hueco, caja], rowHeights=[0.6 * cm, 1.4 * cm, 0.5 * cm])
    firmas.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('TEXTCOLOR', (0, 2), (-1, 2), TENUE), ('FONTSIZE', (0, 2), (-1, 2), 6.5), ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('LINEBELOW', (0, 1), (0, 1), 0.6, TINTA), ('LINEBELOW', (2, 1), (2, 1), 0.6, TINTA), ('LINEBELOW', (4, 1), (4, 1), 0.6, TINTA),
    ]))
    h.append(KeepTogether([
        Spacer(1, 14), firmas, Spacer(1, 8),
        Paragraph(f"Huella SHA-256 de los datos: <font name='Courier'>{datos['huella']}</font>. "
                  'Dos informes con la misma huella contienen exactamente los mismos datos y correcciones.', e['nota']),
    ]))

    doc.build(h, canvasmaker=lambda *a, **k: _LienzoNumerado(*a, encabezado=encabezado_y_pie, **k))
    return salida.getvalue()


# ==============================================================================
# EXCEL
# ==============================================================================

NEGRITA = Font(bold=True)
ENCABEZADO = PatternFill('solid', fgColor='DCE8F8')
RELLENO_FUERA = PatternFill('solid', fgColor='FDE2E2')
RELLENO_ALERTA = PatternFill('solid', fgColor='FEF3C7')
BORDE = Border(bottom=Side(style='thin', color='2A78D6'))


def _hoja_tabla(hoja, encabezados, filas, anchos=None):
    hoja.append(encabezados)
    for celda in hoja[hoja.max_row]:
        celda.font, celda.fill, celda.border = NEGRITA, ENCABEZADO, BORDE
        celda.alignment = Alignment(wrap_text=True, vertical='center')
    inicio = hoja.max_row
    for fila in filas:
        hoja.append(fila)
    hoja.freeze_panes = hoja.cell(row=inicio + 1, column=1)
    if filas:
        hoja.auto_filter.ref = f"A{inicio}:{get_column_letter(len(encabezados))}{hoja.max_row}"
    for k, ancho in enumerate(anchos or [], start=1):
        hoja.column_dimensions[get_column_letter(k)].width = ancho
    return inicio


def _fecha_excel(momento):
    return timezone.localtime(momento).replace(tzinfo=None) if momento else None


def _libro_a_bytes(libro):
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _nombre_hoja(texto, usados):
    """Nombre válido de hoja de Excel (31 caracteres, sin []:*?/\\) y sin repetir."""
    import re
    base = re.sub(r'[\[\]:*?/\\]', '-', texto)[:31]
    nombre, n = base, 2
    while nombre.lower() in usados:
        nombre = f"{base[:28]} {n}"
        n += 1
    usados.add(nombre.lower())
    return nombre


def _hoja_datos(hoja, etapa):
    """Tomas de una etapa: una fila por toma, una columna por parámetro (fuera de límites resaltados)."""
    columnas = etapa['columnas']
    encabezados = ['Tarea', 'Toma', 'Edad lote (h)', 'Edad etapa (h)', 'Adicional', 'Muestras', 'Registró'] + \
                  [f"{c['nombre']} ({c['unidad']})" if c['unidad'] else c['nombre'] for c in columnas]
    filas = []
    for tm in etapa['tomas']:
        fila = [tm['tarea'], _fecha_excel(tm['momento']), tm['edad_lote_h'], tm['edad_etapa_h'],
                'Sí' if tm['adicional'] else '', tm['muestras'], tm['registrado_por']]
        for c in columnas:
            v = tm['valores'].get(c['parametro'])
            fila.append(None if not v else (float(v['numero']) if v['numero'] is not None else v['texto']))
        filas.append(fila)
    inicio = _hoja_tabla(hoja, encabezados, filas, [40, 17, 12, 12, 10, 18, 26] + [14] * len(columnas))
    for k, tm in enumerate(etapa['tomas'], start=inicio + 1):
        hoja.cell(row=k, column=2).number_format = 'dd/mm/yyyy hh:mm'
        for j, c in enumerate(columnas, start=8):
            v = tm['valores'].get(c['parametro'])
            if not v:
                continue
            celda = hoja.cell(row=k, column=j)
            if v['numero'] is not None:
                celda.number_format = '0' + ('.' + '0' * c['decimales'] if c['decimales'] else '')
            if v['fuera_especificacion']:
                celda.fill = RELLENO_FUERA
                celda.font = Font(bold=True, color='B91C1C')
            elif v['fuera_alerta']:
                celda.fill = RELLENO_ALERTA
            if v['corregido']:
                celda.comment = Comment('Valor corregido: ver la hoja «Correcciones».', 'MES')


def _hoja_combinadas(hoja, datos):
    """Datos de las gráficas combinadas y una gráfica nativa de Excel por parámetro, apiladas con el mismo eje X."""
    hoja.append(['Gráficas combinadas: cada parámetro contra la edad del lote (h); todas con el mismo eje X.'])
    hoja['A1'].font = Font(bold=True, size=12, color='2A78D6')
    combinadas = datos['combinadas']
    xs = [p[0] for c in combinadas for p in c['puntos']]
    x_min, x_max = min(0, min(xs)), max(xs)
    x_max = max(x_max + max(2, (x_max - x_min) * 0.04), 24)
    fila_inicio = 3
    for k, c in enumerate(combinadas):
        col = c['columna']
        columna = 1 + k * 3
        hoja.cell(row=fila_inicio, column=columna, value='Edad lote (h)').font = NEGRITA
        titulo = f"{col['nombre']} ({col['unidad']}) · {c['etapa']}" if col['unidad'] else f"{col['nombre']} · {c['etapa']}"
        hoja.cell(row=fila_inicio, column=columna + 1, value=titulo).font = NEGRITA
        for j, (x, y, _) in enumerate(c['puntos'], start=fila_inicio + 1):
            hoja.cell(row=j, column=columna, value=x)
            hoja.cell(row=j, column=columna + 1, value=y)
        hoja.column_dimensions[get_column_letter(columna)].width = 12
        hoja.column_dimensions[get_column_letter(columna + 1)].width = 22

        grafica = ScatterChart()
        grafica.title = titulo
        grafica.style = 13
        grafica.height, grafica.width = 6.5, 24
        grafica.x_axis.title = 'Edad del lote (h)'
        grafica.y_axis.title = col['unidad'] or col['nombre']
        grafica.x_axis.scaling.min, grafica.x_axis.scaling.max = x_min, x_max
        grafica.x_axis.delete = False
        grafica.y_axis.delete = False
        grafica.legend = None
        fin = fila_inicio + len(c['puntos'])
        serie = Series(Reference(hoja, min_col=columna + 1, min_row=fila_inicio + 1, max_row=fin),
                       Reference(hoja, min_col=columna, min_row=fila_inicio + 1, max_row=fin), title=col['nombre'])
        serie.marker.symbol = 'circle'
        serie.marker.size = 6
        color = COLORES_COMBINADAS[k % len(COLORES_COMBINADAS)].lstrip('#')
        serie.graphicalProperties.line.solidFill = color
        serie.graphicalProperties.line.width = 22000
        serie.marker.graphicalProperties.solidFill = color
        serie.marker.graphicalProperties.line.solidFill = color
        grafica.series.append(serie)
        hoja.add_chart(grafica, f"{get_column_letter(len(combinadas) * 3 + 2)}{3 + k * 14}")


def _resumen_retrasos_excel(r):
    """Filas del Resumen con el retraso del lote y su origen principal."""
    if not r:
        return []
    o = r['origen']
    return [
        ('Retraso vs. planeado (h)' + ('' if r['lote']['terminado'] else ', estimado'), r['retraso_horas']),
        ('Origen del retraso', f"[{o['paso']}] {o['nombre']} · {'inicio tarde' if o['tipo'] == 'INICIO' else 'duró más'} "
                               f"{_duracion(o['horas'])}" if o else '—'),
    ]


def _hojas_retrasos(libro, r):
    """Hojas «Retrasos» (todas las actividades, en horas) y «Muestreos tarde». No entran en la huella."""
    ahora = r['ahora']
    estados = {'PENDIENTE': 'Pendiente', 'EN_PROGRESO': 'En curso', 'COMPLETADA': 'Completada', 'OMITIDA': 'Omitida'}
    motivos = lambda a: '\n'.join(
        ([f"Reprogramada: {a['reprogramacion']['motivo']} {a['reprogramacion']['notas']}".strip()] if a.get('reprogramacion') else [])
        + [f"{m['que']} ({_local(m['momento'])}, {m['actor']}): {m['texto']}" for m in a['motivos']]
        + [f"{ev['folio']} {ev['tipo']}: {ev['titulo']}" + (' (abierto)' if ev['estado'] == 'ABIERTO' else '') for ev in a['eventos']])
    hoja = libro.create_sheet('Retrasos')
    hoja.append([f"Retraso {'al terminar' if r['lote']['terminado'] else f'estimado al {_local(ahora)}'}: {r['retraso_horas']} h. "
                 'Horas positivas = retraso; negativas = adelanto. Calculado al emitir: no forma parte de la huella.'])
    hoja['A1'].font = Font(italic=True, color='64748B')
    _hoja_tabla(hoja, ['Paso', 'Actividad', 'Etapa', 'Estado', 'En la ruta', 'Inicio planeado', 'Inicio real/estimado', 'Inicio real',
                       'Fin planeado', 'Fin real/estimado', 'Fin real', 'Heredado (h)', 'Nació al iniciar (h)', 'Nació por duración (h)',
                       'Al terminar (h)', 'Qué pasó', 'Motivos y eventos'],
                [[a['paso'], a['nombre'], a['etapa'], estados.get(a['estado'], a['estado']), 'Sí' if a['en_ruta'] else '',
                  _fecha_excel(a['base_inicio']), _fecha_excel(a['inicio']), 'Sí' if a['inicio_real'] else 'Estimado',
                  _fecha_excel(a['base_fin']), _fecha_excel(a['fin']), 'Sí' if a['fin_real'] else 'Estimado',
                  a['heredado_h'], a['espera_h'], a['duracion_extra_h'], a['retraso_fin_h'],
                  ' · '.join(t for t in (_texto_espera(a, ahora), _texto_duracion(a)) if t), motivos(a)] for a in r['actividades']],
                [8, 32, 16, 12, 9, 17, 17, 10, 17, 17, 10, 12, 14, 16, 13, 40, 60])
    hoja = libro.create_sheet('Muestreos tarde')
    _hoja_tabla(hoja, ['Paso', 'Muestreo', 'Etapa', 'Le tocaba', 'Se tomó', 'Retraso (h)', 'Tolerancia (h)', 'Fuera de tolerancia',
                       'Motivos y eventos'],
                [[m['paso'], m['nombre'], m['etapa'], _fecha_excel(m['esperado']), _fecha_excel(m['inicio']) if m['inicio_real'] else 'Aún no se toma',
                  m['retraso_h'], m['tolerancia_h'], 'Sí' if m['fuera_tolerancia'] else 'No', motivos(m)] for m in r['muestreos']],
                [8, 32, 16, 17, 17, 12, 13, 12, 60])


def excel_lote(datos, emitido_por):
    """Excel del informe de un lote (o de una etapa)."""
    lote = datos['lote']
    libro = Workbook()
    resumen = libro.active
    resumen.title = 'Resumen'
    resumen.append([_titulo_documento(datos)])
    resumen['A1'].font = Font(bold=True, size=14, color='0F172A')
    cero = datos['edad_cero_lote']
    for etiqueta, valor in [
        ('Lote', lote.identificador_lote), ('Receta', datos['receta']), ('Estado', lote.get_estado_display()),
        ('Alcance', f"Etapa {datos['etapa']}" if datos['etapa'] else 'Todo el lote'),
        ('Inicio real', _fecha_excel(datos['inicio_real'])), ('Fin real', _fecha_excel(datos['fin_real'])),
        ('Edad 0 del lote', _fecha_excel(cero['momento']) if cero else None), ('Origen de la edad 0', cero['origen'] if cero else ''),
        ('Eventos GMP', f"{len(datos['eventos'])} ({sum(1 for ev in datos['eventos'] if ev['estado'] == 'Abierto')} abiertos)"),
        *_resumen_retrasos_excel(datos.get('retrasos')),
        ('Emitido por', emitido_por), ('Emitido el', _fecha_excel(datos['generado'])),
        ('Huella SHA-256 de los datos', datos['huella']),
    ]:
        resumen.append([etiqueta, valor])
        resumen.cell(row=resumen.max_row, column=1).font = NEGRITA
        if hasattr(valor, 'year'):
            resumen.cell(row=resumen.max_row, column=2).number_format = 'dd/mm/yyyy hh:mm'
    resumen.append([])
    if datos['indicadores']:
        resumen.append(['Indicadores'])
        resumen.cell(row=resumen.max_row, column=1).font = Font(bold=True, size=12, color='2A78D6')
        _hoja_tabla(resumen, ['Indicador', 'Etapa', 'Tipo', 'Valor', 'Unidad', 'Meta mín', 'Meta máx', 'Resultado', 'Nota'],
                    [[i['nombre'], i['etapa'], i['tipo'], i['valor'], i['unidad'], i['meta_min'], i['meta_max'], i['estado'], i['nota']]
                     for i in datos['indicadores']])
        resumen.freeze_panes = None
    for letra, ancho in zip('ABCDEFGHI', (30, 34, 26, 12, 12, 12, 12, 14, 40)):
        resumen.column_dimensions[letra].width = ancho

    # Datos: una hoja por etapa (como en el diseñador), una fila por toma y una columna por parámetro
    usados = set()
    for etapa in datos['etapas']:
        nombre = _nombre_hoja(f"Datos {etapa['etapa']}", usados)
        _hoja_datos(libro.create_sheet(nombre), etapa)

    if datos['combinadas']:
        _hoja_combinadas(libro.create_sheet(_nombre_hoja('Gráficas combinadas', usados)), datos)

    hoja = libro.create_sheet('Fuera de límites')
    _hoja_tabla(hoja, ['Etapa', 'Parámetro', 'Valor', 'Unidad', 'Límite', 'Rango permitido', 'Tarea', 'Toma', 'Edad lote (h)'],
                [[f['etapa'], f['parametro'], f['valor'], f['unidad'], f['tipo'], f['rango'], f['tarea'], _fecha_excel(f['momento']),
                  f['edad_lote_h']] for f in datos['fuera']], [14, 26, 12, 10, 16, 18, 40, 17, 12])

    hoja = libro.create_sheet('Eventos GMP')
    _hoja_tabla(hoja, ['Folio', 'Ocurrió', 'Tipo', 'Origen', 'Tarea', 'Equipo', 'Título', 'Descripción', 'Acción inmediata', 'Reportó',
                       'Estado', 'Clasificación', 'Nº desviación', 'Evaluación', 'Firmante', 'Significado de la firma', 'Fecha de la firma'],
                [[ev['folio'], _fecha_excel(ev['ocurrio_en']), ev['tipo'], ev['origen'], ev['tarea'], ev['equipo'], ev['titulo'],
                  ev['descripcion'], ev['accion_inmediata'], ev['reporto'], ev['estado'], ev['clasificacion'], ev['numero_desviacion'],
                  ev['evaluacion'], (ev['firma'] or {}).get('nombre', ''), (ev['firma'] or {}).get('significado', ''),
                  _fecha_excel(ev['firma']['fecha']) if ev['firma'] else None] for ev in datos['eventos']],
                [14, 17, 26, 12, 34, 22, 36, 60, 40, 24, 11, 22, 16, 50, 22, 30, 17])

    if datos.get('retrasos'):
        _hojas_retrasos(libro, datos['retrasos'])

    hoja = libro.create_sheet('Muestras adicionales')
    _hoja_tabla(hoja, ['Tarea', 'Toma', 'Código', 'Motivo', 'Registró'],
                [[a['tarea'], _fecha_excel(a['momento']), a['muestras'], a['motivo'], a['registrado_por']] for a in datos['adicionales']],
                [40, 17, 16, 60, 26])

    hoja = libro.create_sheet('Correcciones')
    filas = []
    for c in datos['correcciones']:
        firma = c['firma'] or {}
        for x in c['cambios'] or [{'codigo': '', 'antes': '', 'despues': ''}]:
            filas.append([_fecha_excel(c['fecha']), c['tarea'], x['codigo'], x['antes'], x['despues'], c['motivo'], c['actor'],
                          firma.get('nombre', ''), firma.get('significado', ''), _fecha_excel(firma['fecha']) if firma else None,
                          firma.get('metodo', '')])
    _hoja_tabla(hoja, ['Fecha', 'Tarea', 'Parámetro', 'Antes', 'Después', 'Motivo', 'Corrigió', 'Firmante', 'Significado de la firma',
                       'Fecha de la firma', 'Método'], filas, [17, 36, 14, 10, 10, 44, 22, 22, 34, 17, 14])

    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if hasattr(celda.value, 'year') and celda.number_format == 'General':
                    celda.number_format = 'dd/mm/yyyy hh:mm'
    return _libro_a_bytes(libro)


def excel_lotes(proceso, etapa_nombre, filas, indicadores_por_lote, indicadores, lotes, emitido_por):
    """Excel de varios lotes: Datos en formato largo (una fila por valor) e Indicadores por lote."""
    libro = Workbook()
    info = libro.active
    info.title = 'Información'
    for etiqueta, valor in [
        ('Receta', proceso.nombre), ('Alcance', f"Etapa {etapa_nombre}" if etapa_nombre else 'Todo el lote'),
        ('Lotes', len(lotes)), ('Valores', len(filas)), ('Emitido por', emitido_por), ('Emitido el', _fecha_excel(timezone.now())),
        ('Nota', 'Solo valores vigentes (los corregidos muestran el valor actual; el historial está en el informe de cada lote).'),
    ]:
        info.append([etiqueta, valor])
        info.cell(row=info.max_row, column=1).font = NEGRITA
    info['B6'].number_format = 'dd/mm/yyyy hh:mm'
    info.column_dimensions['A'].width, info.column_dimensions['B'].width = 16, 90

    hoja = libro.create_sheet('Datos')
    inicio = _hoja_tabla(hoja, ['Lote', 'Estado', 'Etapa', 'Tarea', 'Toma', 'Edad lote (h)', 'Edad etapa (h)', 'Adicional', 'Muestras',
                                'Código', 'Parámetro', 'Unidad', 'Valor', 'Texto', 'Calculado', 'Fuera de esp.', 'Fuera de alerta',
                                'Corregido', 'Registró'],
                         [[f['lote'], f['estado'], f['etapa'], f['tarea'], _fecha_excel(f['momento']), f['edad_lote_h'], f['edad_etapa_h'],
                           'Sí' if f['adicional'] else '', f['muestras'], f['codigo'], f['parametro'], f['unidad'],
                           float(f['numero']) if f['numero'] is not None else None, f['texto'], 'Sí' if f['calculado'] else '',
                           'Sí' if f['fuera_especificacion'] else '', 'Sí' if f['fuera_alerta'] else '', 'Sí' if f['corregido'] else '',
                           f['registrado_por']] for f in filas],
                         [14, 14, 12, 34, 17, 11, 11, 9, 14, 14, 22, 10, 11, 14, 10, 11, 11, 10, 22])
    for k in range(inicio + 1, hoja.max_row + 1):
        hoja.cell(row=k, column=5).number_format = 'dd/mm/yyyy hh:mm'
        if hoja.cell(row=k, column=16).value == 'Sí':
            hoja.cell(row=k, column=13).fill = RELLENO_FUERA
        elif hoja.cell(row=k, column=17).value == 'Sí':
            hoja.cell(row=k, column=13).fill = RELLENO_ALERTA

    if indicadores:
        hoja = libro.create_sheet('Indicadores')
        _hoja_tabla(hoja, ['Lote', 'Estado'] + [f"{i.nombre} ({i.unidad})" if i.unidad else i.nombre for i in indicadores],
                    [[l.identificador_lote, l.get_estado_display()] + [indicadores_por_lote.get(l.id, {}).get(i.codigo, {}).get('valor')
                                                                       for i in indicadores] for l in lotes],
                    [14, 14] + [16] * len(indicadores))
    return _libro_a_bytes(libro)
