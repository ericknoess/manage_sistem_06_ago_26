# datos_proceso/tests.py

from datetime import timedelta
from decimal import Decimal

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from auditoria.models import FirmaElectronica, RegistroAuditoria
from procesos.models import AsignacionFaseOperador, EtapaProceso, FaseLote, LoteProduccion, OperacionProceso, ProcesoMaestro
from roster.models import Cuadrilla, Operador, RolOperador
from .formulas import ErrorFormula, evaluar, variables
from .models import IndicadorProceso, ParametroProceso, SeccionCaptura, TomaDatos, ValorParametro
from .services import edades_cero


def usuario_con_rol(nombre, rol):
    usuario = User.objects.create_user(nombre, password='clave-de-prueba')
    usuario.groups.add(Group.objects.get(name=rol))
    return usuario


class FormulasTests(TestCase):

    def test_variables_y_calculo(self):
        self.assertEqual(variables("flujo_total * mezcla_o2 / 100"), ['flujo_total', 'mezcla_o2'])
        self.assertAlmostEqual(evaluar("flujo_total * mezcla_o2 / 100", {'flujo_total': 2, 'mezcla_o2': 25}), 0.5)
        self.assertEqual(evaluar("round(sqrt(x), 1)", {'x': 2}), 1.4)

    def test_falta_un_dato_o_division_entre_cero_da_vacio(self):
        self.assertIsNone(evaluar("a / b", {'a': 1}))
        self.assertIsNone(evaluar("a / b", {'a': 1, 'b': 0}))
        self.assertIsNone(evaluar("log(a)", {'a': -1}))

    def test_rechaza_lo_que_no_es_un_calculo(self):
        for formula in ("__import__('os').system('ls')", "a.b", "lambda: 1", "[1, 2]", "a if b else c", "'texto'", "a == b"):
            with self.subTest(formula=formula), self.assertRaises(ErrorFormula):
                variables(formula)


class CatalogoParametrosTests(TestCase):

    def setUp(self):
        self.api = APIClient()
        self.api.force_authenticate(usuario_con_rol('ingeniera', 'Ingeniería de procesos'))

    def crear(self, **datos):
        return self.api.post('/api/parametros-proceso/', {'tipo_dato': 'NUMERO', 'decimales': 2, **datos}, format='json')

    def test_alta_de_parametro_y_calculado(self):
        self.assertEqual(self.crear(codigo='rot_l', nombre='Rotámetro L', unidad='mm').status_code, 201)
        r = self.crear(codigo='flujo_total', nombre='Flujo total', unidad='L/min', formula='0.0139 * rot_l + 0.059')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data['variables'], ['rot_l'])
        self.assertTrue(RegistroAuditoria.objects.filter(categoria='CATALOGO', accion='CREAR').exists())

    def test_codigo_invalido_o_inmutable(self):
        self.assertEqual(self.crear(codigo='pH Externo', nombre='pH').status_code, 400)
        self.assertEqual(self.crear(codigo='round', nombre='Redondeo').status_code, 400)
        r = self.crear(codigo='ph', nombre='pH')
        cambio = self.api.patch(f"/api/parametros-proceso/{r.data['id']}/", {'codigo': 'ph2'}, format='json')
        self.assertEqual(cambio.status_code, 400)

    def test_formula_con_parametro_inexistente_no_numerico_o_circular(self):
        self.crear(codigo='a', nombre='A')
        self.crear(codigo='aspecto', nombre='Aspecto', tipo_dato='OPCION', opciones=['Normal', 'Turbio'])
        self.assertIn('No existe', str(self.crear(codigo='b', nombre='B', formula='a + zz').data))
        self.assertIn('no es numérico', str(self.crear(codigo='c', nombre='C', formula='aspecto * 2').data))
        b = self.crear(codigo='b', nombre='B', formula='a * 2')
        self.assertEqual(b.status_code, 201)
        self.crear(codigo='c', nombre='C', formula='b + 1')
        # a no puede pasar a depender de c (c → b → a)
        a = ParametroProceso.objects.get(codigo='a')
        r = self.api.patch(f'/api/parametros-proceso/{a.id}/', {'formula': 'c / 2'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('circular', str(r.data))

    def test_opciones_requieren_al_menos_dos(self):
        r = self.crear(codigo='aspecto', nombre='Aspecto', tipo_dato='OPCION', opciones=['Normal'])
        self.assertEqual(r.status_code, 400)

    def test_rol_sin_permiso_no_modifica_el_catalogo(self):
        consulta = APIClient()
        consulta.force_authenticate(usuario_con_rol('consulta', 'Consulta'))
        r = consulta.post('/api/parametros-proceso/', {'codigo': 'ph', 'nombre': 'pH'}, format='json')
        self.assertEqual(r.status_code, 403)
        self.assertEqual(consulta.get('/api/parametros-proceso/').status_code, 200)


class CapturaBase(TestCase):
    """
    Receta (etapa PO4):  T1 Inoculación (1 h, define la edad 0 de PO4 y del lote al completar)
                         → T2 Incubación (24 h, muestreo cada 6 h con muestreo inicial, captura en muestreos)
                         → T3 Cuenta final (1 h, captura al completar)
    Secciones de T2: Reactor (todos), Metabolitos (cada 12 h, muestra PO41-M{n}), Cuenta (a las 0 y 24 h).
    """

    def setUp(self):
        self.admin = APIClient()
        self.admin.force_authenticate(usuario_con_rol('admin', 'Administrador'))
        crear = lambda codigo, nombre, **extra: ParametroProceso.objects.create(codigo=codigo, nombre=nombre, **extra)
        self.temp = crear('temp', 'Temperatura', unidad='°C', decimales=1)
        self.rot_l = crear('rot_l', 'Rotámetro L', decimales=1)
        self.mezcla_o2 = crear('mezcla_o2', 'Mezcla O2', unidad='%', decimales=1)
        self.flujo_total = crear('flujo_total', 'Flujo total', unidad='L/min', decimales=3, formula='0.0139 * rot_l + 0.059')
        self.flujo_o2 = crear('flujo_o2', 'Flujo O2', unidad='L/min', decimales=3, formula='flujo_total * mezcla_o2 / 100')
        self.aspecto = crear('aspecto', 'Aspecto', tipo_dato='OPCION', opciones=['Normal', 'Espuma'])
        self.glucosa = crear('glucosa', 'Glucosa', unidad='g/L')
        self.celulas = crear('celulas', 'Células', unidad='MCel/mL')
        self.viabilidad = crear('viabilidad', 'Viabilidad', unidad='%', decimales=1)

        self.proceso = ProcesoMaestro.objects.create(nombre='EPO prueba')
        self.etapa = EtapaProceso.objects.create(proceso=self.proceso, nombre='PO4', orden=1)
        self.t1 = self.crear_operacion({
            'identificador_paso': 'T1', 'nombre': 'Inoculación', 'tipo_operacion': 'ACTIVA', 'duracion_horas': 1,
            'edad_cero': 'FIN', 'edad_cero_lote': True,
        })
        self.t2 = self.crear_operacion({
            'identificador_paso': 'T2', 'nombre': 'Incubación', 'tipo_operacion': 'INCUBACION', 'duracion_horas': 24,
            'frecuencia_muestreo_horas': 6, 'duracion_muestreo_horas': 0.5, 'ops_muestreo': 1, 'muestreo_inicial': True,
            'predecesora': self.t1.id, 'captura_datos': 'MUESTREOS', 'secciones_captura': [
                {'nombre': 'Reactor', 'frecuencia': 'TODOS', 'parametros': [
                    self.p(self.temp, especificacion_min='36', especificacion_max='38', alerta_min='36.5', alerta_max='37.5'),
                    self.p(self.rot_l), self.p(self.mezcla_o2), self.p(self.flujo_total), self.p(self.flujo_o2),
                    self.p(self.aspecto, obligatorio=False)]},
                {'nombre': 'Metabolitos', 'frecuencia': 'CADA_N', 'cada_horas': 12, 'muestra_fisica': True,
                 'patron_codigo_muestra': 'PO41-M{n}', 'parametros': [self.p(self.glucosa)]},
                {'nombre': 'Cuenta', 'frecuencia': 'HORAS', 'horas': [0, 24], 'parametros': [self.p(self.celulas), self.p(self.viabilidad)]},
            ],
        })
        self.t3 = self.crear_operacion({
            'identificador_paso': 'T3', 'nombre': 'Cuenta final', 'tipo_operacion': 'ACTIVA', 'duracion_horas': 1,
            'predecesora': self.t2.id, 'captura_datos': 'COMPLETAR', 'secciones_captura': [
                {'nombre': 'Cuenta', 'parametros': [self.p(self.celulas)]},
            ],
        })

        self.rol = RolOperador.objects.create(nombre='Operador base')
        cuadrilla = Cuadrilla.objects.create(identificador='A', nombre='Cuadrilla A')
        self.operador = Operador.objects.create(cuadrilla=cuadrilla, nombre='Ana', codigo_empleado='OP-1', rol=self.rol, pin_acceso='1234')

        inicio = timezone.localtime(timezone.now()).replace(second=0, microsecond=0, tzinfo=None)
        r = self.admin.post('/api/lotes/', {'identificador_lote': 'L-1', 'proceso_maestro': self.proceso.id,
                                            'fecha_inicio_planeada': inicio.isoformat()}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.lote = LoteProduccion.objects.get(identificador_lote='L-1')

    @staticmethod
    def p(parametro, **extra):
        return {'parametro': parametro.id, 'obligatorio': True, **extra}

    def crear_operacion(self, datos, esperado=201):
        r = self.admin.post('/api/operaciones-proceso/', {'proceso': self.proceso.id, 'etapa': self.etapa.id, **datos}, format='json')
        self.assertEqual(r.status_code, esperado, r.data)
        return OperacionProceso.objects.get(pk=r.data['id']) if esperado == 201 else r

    def muestreo(self, indice):
        return FaseLote.objects.get(lote=self.lote, operacion_maestra=self.t2, es_subtarea_muestreo=True, indice_muestreo=indice)

    def actividad(self, operacion):
        return FaseLote.objects.get(lote=self.lote, operacion_maestra=operacion, es_subtarea_muestreo=False)

    def forzar(self, fase, estado, momento=None):
        """Cambia el estado sin pasar por la API (para preparar escenarios)."""
        momento = momento or timezone.now()
        fase.estado = estado
        if estado in ('EN_PROGRESO', 'COMPLETADA') and not fase.fecha_inicio_real:
            fase.fecha_inicio_real = momento
        if estado == 'COMPLETADA':
            fase.fecha_fin_real = momento
        fase.save()
        return fase

    def valores_reactor(self, temp='37.0'):
        return {self.temp.id: temp, self.rot_l.id: '100', self.mezcla_o2.id: '20'}

    def registrar(self, cliente, fase, valores, **extra):
        return cliente.post(f'/api/captura/{fase.id}/registrar/', {'valores': valores, **extra}, format='json')


class RecetaYPlanTests(CapturaBase):

    def test_plan_de_cada_muestreo_segun_sus_secciones(self):
        horas = {m.indice_muestreo * 6: m for m in FaseLote.objects.filter(lote=self.lote, operacion_maestra=self.t2, es_subtarea_muestreo=True)}
        self.assertEqual(sorted(horas), [0, 6, 12, 18, 24])  # incluye el muestreo inicial M-0h
        secciones = lambda h: [s['nombre'] for s in horas[h].plan_captura['secciones']]
        self.assertEqual(secciones(0), ['Reactor', 'Metabolitos', 'Cuenta'])
        self.assertEqual(secciones(6), ['Reactor'])
        self.assertEqual(secciones(12), ['Reactor', 'Metabolitos'])
        self.assertEqual(secciones(24), ['Reactor', 'Metabolitos', 'Cuenta'])
        codigos = [next(s['codigo_muestra'] for s in horas[h].plan_captura['secciones'] if s['nombre'] == 'Metabolitos') for h in (0, 12, 24)]
        self.assertEqual(codigos, ['PO41-M01', 'PO41-M02', 'PO41-M03'])
        # La incubación misma no captura (lo hacen sus muestreos); la cuenta final sí, al completar
        self.assertEqual(self.actividad(self.t2).plan_captura, {})
        self.assertEqual([s['nombre'] for s in self.actividad(self.t3).plan_captura['secciones']], ['Cuenta'])

    def test_los_planes_guardan_cada_lista_de_parametros_una_sola_vez(self):
        from .models import PlanParametros
        from .planes import expandir_plan
        m6 = self.muestreo(1)
        self.assertNotIn('parametros', m6.plan_captura['secciones'][0])           # en la base va la huella
        self.assertEqual(len(m6.plan_captura['secciones'][0]['parametros_ref']), 64)
        # 5 muestreos + la cuenta final comparten 4 listas distintas (Reactor, Metabolitos, Cuenta de T2 y de T3)
        self.assertEqual(PlanParametros.objects.count(), 4)
        completo = expandir_plan(m6.plan_captura)
        self.assertEqual([p['codigo'] for p in completo['secciones'][0]['parametros']][:2], ['temp', 'rot_l'])
        lista = PlanParametros.objects.first()
        with self.assertRaises(PermissionError):
            lista.save()
        with self.assertRaises(PermissionError):
            lista.delete()

    def test_editar_la_receta_no_cambia_un_lote_en_curso(self):
        SeccionCaptura.objects.filter(operacion=self.t2, nombre='Metabolitos').delete()
        self.assertIn('Metabolitos', [s['nombre'] for s in self.muestreo(2).plan_captura['secciones']])

    def test_reglas_del_disenador(self):
        base = {'identificador_paso': 'T9', 'nombre': 'Otra', 'tipo_operacion': 'INCUBACION', 'duracion_horas': 24,
                'frecuencia_muestreo_horas': 6, 'captura_datos': 'MUESTREOS'}
        seccion = lambda **extra: [{'nombre': 'S', 'parametros': [self.p(self.glucosa)], **extra}]
        casos = [
            ({**base, 'secciones_captura': seccion(frecuencia='CADA_N', cada_horas=8)}, 'múltiplo'),
            ({**base, 'secciones_captura': seccion(frecuencia='HORAS', horas=[0, 7])}, 'no hay muestreo'),
            ({**base, 'tipo_operacion': 'ACTIVA', 'frecuencia_muestreo_horas': 0, 'secciones_captura': seccion()}, 'incubación'),
            ({**base, 'secciones_captura': []}, 'al menos una sección'),
            ({**base, 'secciones_captura': [{'nombre': 'A', 'parametros': [self.p(self.glucosa)]},
                                            {'nombre': 'B', 'parametros': [self.p(self.glucosa)]}]}, 'déjalo solo en una'),
            ({**base, 'captura_datos': 'NINGUNA', 'edad_cero': 'INICIO'}, 'ya la define'),
            ({**base, 'captura_datos': 'NINGUNA', 'edad_cero_lote': True}, 'debe definir la edad 0'),
        ]
        for datos, mensaje in casos:
            with self.subTest(mensaje=mensaje):
                r = self.crear_operacion(datos, esperado=400)
                self.assertIn(mensaje, str(r.data))

    def test_clonar_receta_copia_captura_y_edad_cero(self):
        r = self.admin.post(f'/api/procesos-maestros/{self.proceso.id}/clonar/', {'nuevo_nombre': 'EPO copia'}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        copia = OperacionProceso.objects.get(proceso_id=r.data['nuevo_id'], identificador_paso='T2')
        self.assertEqual(copia.captura_datos, 'MUESTREOS')
        self.assertTrue(copia.muestreo_inicial)
        self.assertEqual([s.nombre for s in copia.secciones_captura.all()], ['Reactor', 'Metabolitos', 'Cuenta'])
        self.assertEqual(copia.secciones_captura.get(nombre='Reactor').parametros.count(), 6)
        self.assertEqual(OperacionProceso.objects.get(proceso_id=r.data['nuevo_id'], identificador_paso='T1').edad_cero, 'FIN')

    def test_valor_esperado_y_curva_de_la_tendencia(self):
        from .services import plan_de_operacion
        base = {'identificador_paso': 'T4', 'nombre': 'Lavado', 'tipo_operacion': 'ACTIVA', 'duracion_horas': 1,
                'captura_datos': 'COMPLETAR'}
        seccion = lambda **extra: [{'nombre': 'Metabolitos', 'parametros': [self.p(self.glucosa, **extra)]}]
        op = self.crear_operacion({**base, 'secciones_captura': seccion(
            objetivo='0.8', eje_edad='ETAPA', curva_referencia=[[24, 2.1], [0, 3.5], ['12', '2,8'.replace(',', '.')]])})
        ps = op.secciones_captura.get().parametros.get()
        self.assertEqual(ps.curva_referencia, [[0.0, 3.5], [12.0, 2.8], [24.0, 2.1]])  # se ordena por edad
        ref = plan_de_operacion(op)['secciones'][0]['parametros'][0]['referencia']
        self.assertEqual(ref, {'objetivo': '0.8', 'curva': [[0.0, 3.5], [12.0, 2.8], [24.0, 2.1]], 'eje_edad': 'ETAPA'})
        casos = [
            (seccion(curva_referencia=[[0, 1], [0, 2]]), 'repetidas'),
            (seccion(curva_referencia=[[-1, 1]]), 'negativas'),
            (seccion(curva_referencia=[[1, 'x']]), 'dos números'),
            ([{'nombre': 'S', 'parametros': [self.p(self.aspecto, objetivo='1')]}], 'no es numérico'),
        ]
        for i, (secciones, mensaje) in enumerate(casos):
            with self.subTest(mensaje=mensaje):
                r = self.crear_operacion({**base, 'identificador_paso': f'T{5 + i}', 'secciones_captura': secciones}, esperado=400)
                self.assertIn(mensaje, str(r.data))


class CapturaEnPlantaTests(CapturaBase):

    def setUp(self):
        super().setUp()
        self.web = APIClient()
        self.web.force_authenticate(usuario_con_rol('supervisora', 'Jefatura'))
        self.hito = timezone.now() - timedelta(hours=12, minutes=30)
        self.forzar(self.actividad(self.t1), 'COMPLETADA', self.hito)
        self.forzar(self.actividad(self.t2), 'EN_PROGRESO', self.hito)
        self.m12 = self.forzar(self.muestreo(2), 'EN_PROGRESO')

    def test_registra_calcula_y_marca_limites(self):
        r = self.registrar(self.web, self.m12, {**self.valores_reactor(temp='38.4'), self.glucosa.id: '4.5', self.aspecto.id: 'Espuma'})
        self.assertEqual(r.status_code, 201, r.data)
        valores = {v.parametro.codigo: v for v in ValorParametro.objects.filter(toma__fase=self.m12)}
        self.assertEqual(valores['flujo_total'].valor_numero, Decimal('1.449'))   # 0.0139 × 100 + 0.059
        self.assertEqual(valores['flujo_o2'].valor_numero, Decimal('0.290'))      # calculado a partir de otro calculado
        self.assertTrue(valores['flujo_o2'].calculado)
        self.assertTrue(valores['temp'].fuera_especificacion)
        self.assertEqual(valores['aspecto'].valor_texto, 'Espuma')
        self.assertEqual(TomaDatos.objects.get(fase=self.m12).codigos_muestra, {'Metabolitos': 'PO41-M02'})
        self.assertTrue(any('fuera de especificación' in a for a in r.data['avisos']))
        self.assertEqual(r.data['toma']['edad_etapa_h'], r.data['toma']['edad_lote_h'])
        self.assertAlmostEqual(r.data['toma']['edad_etapa_h'], 12.5, delta=0.05)  # edad real desde el fin de la inoculación
        self.assertTrue(RegistroAuditoria.objects.filter(categoria='DATOS_PROCESO', accion='REGISTRAR_DATOS', fase=self.m12).exists())

    def test_obligatorios_y_una_sola_toma(self):
        r = self.registrar(self.web, self.m12, self.valores_reactor())
        self.assertEqual(r.status_code, 400)
        self.assertIn('Glucosa', r.data['error'])
        self.assertEqual(self.registrar(self.web, self.m12, {**self.valores_reactor(), self.glucosa.id: '4'}).status_code, 201)
        self.assertEqual(self.registrar(self.web, self.m12, {**self.valores_reactor(), self.glucosa.id: '4'}).status_code, 400)

    def test_valores_no_validos(self):
        r = self.registrar(self.web, self.m12, {**self.valores_reactor(), self.glucosa.id: 'cuatro', self.aspecto.id: 'Rara'})
        self.assertEqual(r.status_code, 400)
        self.assertIn('Glucosa', r.data['error'])
        self.assertIn('Aspecto', r.data['error'])

    def test_no_se_captura_si_la_tarea_no_esta_en_curso(self):
        m18 = self.muestreo(3)
        r = self.registrar(self.web, m18, self.valores_reactor())
        self.assertEqual(r.status_code, 400)
        self.assertIn('iníciala', r.data['error'])

    def test_no_se_completa_sin_los_datos(self):
        cuerpo = {'estado': 'COMPLETADA', 'firma_password': 'clave-de-prueba'}
        AsignacionFaseOperador.objects.create(fase_lote=self.m12, operador=self.operador, rol_ejercido=self.rol)
        r = self.web.post(f'/api/fases-lote/{self.m12.id}/cambiar-estado/', cuerpo, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.data['codigo'], 'DATOS_PENDIENTES')
        self.registrar(self.web, self.m12, {**self.valores_reactor(), self.glucosa.id: '4'})
        r = self.web.post(f'/api/fases-lote/{self.m12.id}/cambiar-estado/', cuerpo, format='json')
        self.assertEqual(r.status_code, 200, r.data)

    def test_vista_previa_calcula_sin_guardar(self):
        r = self.web.post(f'/api/captura/{self.m12.id}/vista-previa/', {'valores': self.valores_reactor()}, format='json')
        self.assertEqual(r.status_code, 200)
        calculado = next(v for v in r.data['valores'] if v['parametro'] == self.flujo_total.id)
        self.assertEqual(calculado['valor'], '1.449')
        self.assertEqual(r.data['faltantes'], ['Glucosa'])
        self.assertFalse(TomaDatos.objects.exists())

    def test_correccion_con_motivo_y_firma_recalcula(self):
        self.registrar(self.web, self.m12, {**self.valores_reactor(), self.glucosa.id: '4'})
        toma = TomaDatos.objects.get(fase=self.m12)
        url = f'/api/captura/{self.m12.id}/corregir/'
        cuerpo = {'toma': toma.id, 'parametro': self.rot_l.id, 'valor': '120'}
        sin_motivo = self.web.post(url, {**cuerpo, 'firma_password': 'clave-de-prueba'}, format='json')
        self.assertEqual(sin_motivo.data['codigo'], 'MOTIVO_REQUERIDO')
        mala_firma = self.web.post(url, {**cuerpo, 'motivo': 'Lectura mal transcrita', 'firma_password': 'otra'}, format='json')
        self.assertEqual(mala_firma.data['codigo'], 'FIRMA_INVALIDA')

        r = self.web.post(url, {**cuerpo, 'motivo': 'Lectura mal transcrita', 'firma_password': 'clave-de-prueba'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        rot = ValorParametro.objects.filter(toma=toma, parametro=self.rot_l)
        self.assertEqual(rot.count(), 2)
        self.assertEqual(rot.get(vigente=True).valor_numero, Decimal('120.0'))
        self.assertEqual(rot.get(vigente=True).reemplaza, rot.get(vigente=False))
        flujo = ValorParametro.objects.get(toma=toma, parametro=self.flujo_total, vigente=True)
        self.assertEqual(flujo.valor_numero, Decimal('1.727'))  # recalculado: 0.0139 × 120 + 0.059
        registro = RegistroAuditoria.objects.get(accion='CORREGIR_DATO')
        self.assertEqual(registro.cambios['rot_l'], ['100', '120'])
        self.assertIn('flujo_total', registro.cambios)
        self.assertTrue(FirmaElectronica.objects.filter(registro=registro).exists())

    def test_no_se_corrige_un_calculado_ni_se_borra_un_dato(self):
        self.registrar(self.web, self.m12, {**self.valores_reactor(), self.glucosa.id: '4'})
        toma = TomaDatos.objects.get(fase=self.m12)
        r = self.web.post(f'/api/captura/{self.m12.id}/corregir/', {
            'toma': toma.id, 'parametro': self.flujo_total.id, 'valor': '2', 'motivo': 'Prueba de corrección', 'firma_password': 'clave-de-prueba',
        }, format='json')
        self.assertEqual(r.status_code, 400)
        with self.assertRaises(PermissionError):
            toma.delete()

    def test_operador_con_pin_solo_en_sus_tareas(self):
        movil = APIClient()
        movil.post('/api/mobile/login/', {'codigo_empleado': 'OP-1', 'pin_acceso': '1234'}, format='json')
        valores = {**self.valores_reactor(), self.glucosa.id: '4'}
        self.assertEqual(self.registrar(movil, self.m12, valores).status_code, 403)
        AsignacionFaseOperador.objects.create(fase_lote=self.m12, operador=self.operador, rol_ejercido=self.rol)
        r = self.registrar(movil, self.m12, valores)
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(TomaDatos.objects.get(fase=self.m12).registrado_por_operador, self.operador)

    def test_rol_sin_ejecutar_fase_no_captura(self):
        consulta = APIClient()
        consulta.force_authenticate(usuario_con_rol('consulta', 'Consulta'))
        self.assertEqual(self.registrar(consulta, self.m12, self.valores_reactor()).status_code, 403)
        self.assertEqual(consulta.get(f'/api/captura/{self.m12.id}/').status_code, 200)

    def test_muestra_adicional_con_motivo(self):
        incubacion = self.actividad(self.t2)
        url = f'/api/captura/{incubacion.id}/adicional/'
        r = self.web.post(url, {'valores': {self.glucosa.id: '3.1'}, 'motivo': 'corto'}, format='json')
        self.assertEqual(r.status_code, 400)
        r = self.web.post(url, {'valores': {self.glucosa.id: '3.1'}, 'motivo': 'Revisión de consumo de glucosa', 'codigo_muestra': 'ADIC-1'}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        toma = TomaDatos.objects.get(fase=incubacion)
        self.assertTrue(toma.adicional)
        self.assertEqual(toma.codigos_muestra, {'Muestra adicional': 'ADIC-1'})
        self.assertEqual(len(r.data['adicionales']), 1)
        self.assertEqual(self.web.post(f'/api/captura/{self.m12.id}/adicional/', {'valores': {self.glucosa.id: '3'}, 'motivo': 'Desde un muestreo no aplica'},
                                       format='json').status_code, 400)

    def test_tendencia_con_los_datos_anteriores_de_la_etapa(self):
        m0 = self.forzar(self.muestreo(0), 'EN_PROGRESO', self.hito)
        r = self.registrar(self.web, m0, {**self.valores_reactor(), self.glucosa.id: '4.2', self.celulas.id: '0.3', self.viabilidad.id: '98'},
                           momento=(self.hito + timedelta(minutes=30)).isoformat())
        self.assertEqual(r.status_code, 201, r.data)
        detalle = self.web.get(f'/api/captura/{self.m12.id}/').data
        glucosa = detalle['tendencias'][str(self.glucosa.id)]
        self.assertEqual([p['valor'] for p in glucosa], [4.2])
        self.assertAlmostEqual(glucosa[0]['edad_lote_h'], 0.5, delta=0.05)
        self.assertEqual(len(detalle['tendencias'][str(self.temp.id)]), 1)
        self.assertNotIn(str(self.aspecto.id), detalle['tendencias'])  # solo numéricos
        self.assertIn('referencia', detalle['plan']['secciones'][0]['parametros'][0])

    def test_tabla_del_lote_con_edades(self):
        self.registrar(self.web, self.m12, {**self.valores_reactor(), self.glucosa.id: '4'})
        r = self.web.get(f'/api/captura/lote/?lote={self.lote.id}')
        self.assertEqual(r.status_code, 200)
        fila = r.data['filas'][0]
        self.assertEqual(fila['valores']['glucosa']['valor'], '4')
        self.assertEqual(fila['etapa'], 'PO4')
        self.assertIn('glucosa', [p['codigo'] for p in r.data['parametros']])

    def test_resumen_en_la_fase_y_en_la_tablet(self):
        AsignacionFaseOperador.objects.create(fase_lote=self.m12, operador=self.operador, rol_ejercido=self.rol)
        fase = self.web.get(f'/api/fases-lote/{self.m12.id}/').data
        self.assertEqual(fase['captura']['secciones'], ['Reactor', 'Metabolitos'])
        self.assertFalse(fase['captura']['capturada'])
        movil = APIClient()
        movil.post('/api/mobile/login/', {'codigo_empleado': 'OP-1', 'pin_acceso': '1234'}, format='json')
        tareas = movil.get(f'/api/mobile/{self.operador.id}/tareas-pendientes/').data['tareas_pendientes']
        self.assertTrue(tareas[0]['captura']['requerida'])


class EdadCeroTests(CapturaBase):

    def test_estimada_antes_del_hito_y_real_despues(self):
        ceros = edades_cero(self.lote)
        self.assertTrue(ceros[self.etapa.id]['estimada'])
        self.assertEqual(ceros[self.etapa.id]['momento'], self.actividad(self.t1).fin_programado)
        fin_real = timezone.now() - timedelta(hours=3)
        self.forzar(self.actividad(self.t1), 'COMPLETADA', fin_real)
        ceros = edades_cero(self.lote)
        self.assertFalse(ceros[self.etapa.id]['estimada'])
        self.assertEqual(ceros[self.etapa.id]['momento'], fin_real)
        self.assertEqual(ceros['lote']['momento'], fin_real)

    def test_sin_actividad_de_edad_cero_usa_el_primer_inicio(self):
        OperacionProceso.objects.filter(pk=self.t1.pk).update(edad_cero='', edad_cero_lote=False)
        ceros = edades_cero(self.lote)
        self.assertIn('no hay actividad de edad 0', ceros[self.etapa.id]['origen'])
        self.assertEqual(ceros[self.etapa.id]['momento'], self.actividad(self.t1).inicio_programado)


class CalculosAnalisisTests(TestCase):

    def test_area_por_trapecios_e_interpolacion(self):
        from .services import area_curva_entre, area_trapecios, interpolar
        self.assertEqual(area_trapecios([(0, 1), (10, 3), (10, 9), (20, 3)]), 80.0)  # 20 + 60: el tramo sin avance de edad no suma
        curva = [[0, 0], [10, 10], [20, 10]]
        self.assertEqual(interpolar(curva, 5), 5)
        self.assertEqual(interpolar(curva, -3), 0)
        self.assertEqual(interpolar(curva, 99), 10)
        self.assertEqual(area_curva_entre(curva, 5, 15), 37.5 + 50)  # 5→10 h sube de 5 a 10; 10→15 h se queda en 10


class AnalisisBase(CapturaBase):
    """Células a las 0 h (0.3) y a las 24 h (0.5) → ABC = 9.6; curva esperada 0.3 → 0.7 → área 12 → apego 80 %."""

    def setUp(self):
        super().setUp()
        from .services import generar_plan_captura
        self.web = APIClient()
        self.web.force_authenticate(usuario_con_rol('ingeniera', 'Ingeniería de procesos'))
        ps = SeccionCaptura.objects.get(operacion=self.t2, nombre='Cuenta').parametros.get(parametro=self.celulas)
        ps.curva_referencia = [[0, 0.3], [24, 0.7]]
        ps.save()
        generar_plan_captura(self.lote)
        hito = timezone.now() - timedelta(hours=30)
        self.forzar(self.actividad(self.t1), 'COMPLETADA', hito)
        self.forzar(self.actividad(self.t2), 'EN_PROGRESO', hito)
        for indice, horas, celulas in ((0, 0, '0.3'), (4, 24, '0.5')):
            fase = self.forzar(self.muestreo(indice), 'EN_PROGRESO', hito)
            r = self.registrar(self.admin, fase, {**self.valores_reactor(), self.glucosa.id: '4', self.celulas.id: celulas,
                                                  self.viabilidad.id: '97'}, momento=(hito + timedelta(hours=horas)).isoformat())
            self.assertEqual(r.status_code, 201, r.data)

    def indicador(self, esperado=201, **datos):
        r = self.web.post('/api/indicadores-proceso/', {'proceso': self.proceso.id, 'decimales': 2, **datos}, format='json')
        self.assertEqual(r.status_code, esperado, r.data)
        return r

    def crear_indicadores(self):
        self.indicador(codigo='abc_cel', nombre='ABC células', tipo='AREA', parametro=self.celulas.id)
        self.indicador(codigo='apego', nombre='Apego', tipo='APEGO', parametro=self.celulas.id, etapa=self.etapa.id, meta_min='90')
        self.indicador(codigo='gramos', nombre='Gramos estimados', tipo='FORMULA', formula='abc_cel * 10 - 5')
        self.indicador(codigo='glu_final', nombre='Glucosa final', tipo='ULTIMO', parametro=self.glucosa.id)



class AnalisisTests(AnalisisBase):

    def test_indicadores_del_lote(self):
        self.crear_indicadores()
        r = self.web.get(f'/api/analisis-proceso/?proceso={self.proceso.id}&lotes={self.lote.id}')
        self.assertEqual(r.status_code, 200, r.data)
        valores = r.data['valores_indicadores'][str(self.lote.id)]
        self.assertEqual(valores['abc_cel']['valor'], 9.6)
        self.assertEqual(valores['apego']['valor'], 80.0)
        self.assertEqual(valores['gramos']['valor'], 91.0)
        self.assertEqual(valores['glu_final']['valor'], 4.0)
        celulas = next(p for p in r.data['parametros'] if p['codigo'] == 'celulas')
        puntos = celulas['series'][str(self.lote.id)]   # [edad_lote, edad_etapa, valor, estado, adicional, momento, tarea]
        self.assertEqual([q[2] for q in puntos], [0.3, 0.5])
        self.assertAlmostEqual(puntos[1][0], 24, delta=0.05)
        self.assertIn('M-24h', r.data['tareas'][puntos[1][6]])
        self.assertEqual(celulas['referencia']['curva'], [[0.0, 0.3], [24.0, 0.7]])
        temp = next(p for p in r.data['parametros'] if p['codigo'] == 'temp')
        self.assertEqual(temp['limites']['especificacion_max'], '38')
        self.assertTrue(RegistroAuditoria.objects.filter(categoria='RECETA', accion='CREAR', descripcion__icontains='indicador').exists())

    def test_sin_curva_o_sin_datos_explica_por_que(self):
        self.indicador(codigo='apego_glu', nombre='Apego glucosa', tipo='APEGO', parametro=self.glucosa.id)
        self.indicador(codigo='abc_temp', nombre='ABC T3', tipo='AREA', parametro=self.temp.id, etapa=EtapaProceso.objects.create(
            proceso=self.proceso, nombre='Producción', orden=2).id)
        valores = self.web.get(f'/api/analisis-proceso/?proceso={self.proceso.id}&lotes={self.lote.id}').data['valores_indicadores'][str(self.lote.id)]
        self.assertIsNone(valores['apego_glu']['valor'])
        self.assertIn('curva esperada', valores['apego_glu']['nota'])
        self.assertEqual(valores['abc_temp']['nota'], 'Sin datos')

    def test_reglas_de_los_indicadores(self):
        self.crear_indicadores()
        otra = ProcesoMaestro.objects.create(nombre='Otra receta')
        casos = [
            ({'codigo': 'abc_cel', 'nombre': 'Repetido', 'tipo': 'AREA', 'parametro': self.celulas.id}, 'Ya hay'),
            ({'codigo': 'g2', 'nombre': 'G2', 'tipo': 'FORMULA', 'formula': 'zz * 2'}, 'No existen'),
            ({'codigo': 'g3', 'nombre': 'G3', 'tipo': 'FORMULA', 'formula': 'g3 + 1'}, 'propio'),
            ({'codigo': 'asp', 'nombre': 'Asp', 'tipo': 'ULTIMO', 'parametro': self.aspecto.id}, 'numérico'),
            ({'codigo': 'sin', 'nombre': 'Sin', 'tipo': 'AREA'}, 'Elige el parámetro'),
            ({'codigo': 'e', 'nombre': 'E', 'tipo': 'AREA', 'parametro': self.celulas.id,
              'etapa': EtapaProceso.objects.create(proceso=otra, nombre='X', orden=1).id}, 'no es de esta receta'),
        ]
        for datos, mensaje in casos:
            with self.subTest(mensaje=mensaje):
                self.assertIn(mensaje, str(self.indicador(esperado=400, **datos).data))
        # circular: abc_cel no puede pasar a calcularse con gramos (gramos → abc_cel)
        abc = IndicadorProceso.objects.get(proceso=self.proceso, codigo='abc_cel')
        r = self.web.patch(f'/api/indicadores-proceso/{abc.id}/', {'tipo': 'FORMULA', 'formula': 'gramos / 2'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('circular', str(r.data))

    def test_operador_no_configura_y_clonar_copia_indicadores(self):
        operador = APIClient()
        operador.force_authenticate(usuario_con_rol('op', 'Operador'))
        r = operador.post('/api/indicadores-proceso/', {'proceso': self.proceso.id, 'codigo': 'x', 'nombre': 'X', 'tipo': 'AREA',
                                                        'parametro': self.celulas.id}, format='json')
        self.assertEqual(r.status_code, 403)
        self.crear_indicadores()
        r = self.admin.post(f'/api/procesos-maestros/{self.proceso.id}/clonar/', {'nuevo_nombre': 'EPO copia'}, format='json')
        copia = IndicadorProceso.objects.get(proceso_id=r.data['nuevo_id'], codigo='apego')
        self.assertNotEqual(copia.etapa_id, self.etapa.id)
        self.assertEqual(copia.etapa.nombre, 'PO4')

    def test_graficas_por_etapa_como_en_el_disenador(self):
        produccion = EtapaProceso.objects.create(proceso=self.proceso, nombre='Producción', orden=2)
        self.crear_operacion({'identificador_paso': 'T9', 'nombre': 'Perfusión', 'tipo_operacion': 'ACTIVA', 'duracion_horas': 1,
                              'etapa': produccion.id, 'captura_datos': 'COMPLETAR',
                              'secciones_captura': [{'nombre': 'Reactor', 'parametros': [self.p(self.celulas, especificacion_min='1')]}]})
        r = self.web.get(f'/api/analisis-proceso/?proceso={self.proceso.id}&lotes={self.lote.id}')
        self.assertEqual([e['nombre'] for e in r.data['etapas']], ['PO4', 'Producción'])
        celulas = [p for p in r.data['parametros'] if p['codigo'] == 'celulas']
        self.assertEqual([p['etapa'] for p in celulas], ['PO4', 'Producción'])       # una gráfica por etapa
        self.assertEqual(celulas[0]['bloque'], '[T2] Incubación · Cuenta')
        self.assertEqual(len(celulas[0]['series'][str(self.lote.id)]), 2)             # los datos de PO4
        self.assertEqual(celulas[1]['series'], {})                                      # Producción aún sin datos
        self.assertEqual(celulas[1]['limites']['especificacion_min'], '1')             # con los límites de esa etapa
        recetas = self.web.get('/api/analisis-proceso/recetas/').data
        receta = next(x for x in recetas if x['id'] == self.proceso.id)
        self.assertEqual((receta['etapas'], receta['lotes'], receta['lotes_con_datos']), (2, 1, 1))

    def test_lotes_de_la_receta_y_limite_de_lotes(self):
        r = self.web.get(f'/api/analisis-proceso/lotes/?proceso={self.proceso.id}')
        self.assertEqual(r.data['lotes'][0]['tomas'], 2)
        self.assertEqual(r.data['etapas'][0]['nombre'], 'PO4')
        muchos = ','.join(str(i) for i in range(1, 202))
        self.assertEqual(self.web.get(f'/api/analisis-proceso/?proceso={self.proceso.id}&lotes={muchos}').status_code, 400)
        repetidos = ','.join(str(self.lote.id) for _ in range(300))   # los repetidos cuentan una vez
        self.assertEqual(len(self.web.get(f'/api/analisis-proceso/?proceso={self.proceso.id}&lotes={repetidos}').data['lotes']), 1)


class InformesTests(AnalisisBase):
    """Informes PDF y Excel sobre los datos de AnalisisTests (células a las 0 y 24 h)."""

    def test_pdf_y_excel_del_lote_quedan_en_la_bitacora(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.crear_indicadores()
        r = self.web.get(f'/api/informes-proceso/{self.lote.id}/pdf/?etapa={self.etapa.id}')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertTrue(r.content.startswith(b'%PDF'))
        self.assertIn('Informe_L-1_PO4_', r['Content-Disposition'])
        r = self.web.get(f'/api/informes-proceso/{self.lote.id}/excel/')
        self.assertEqual(r.status_code, 200)
        libro = load_workbook(BytesIO(r.content))
        self.assertEqual(libro.sheetnames, ['Resumen', 'Datos PO4', 'Fuera de límites', 'Muestras adicionales', 'Correcciones'])
        datos = libro['Datos PO4']
        encabezados = [c.value for c in datos[1]]
        celulas = encabezados.index('Células (MCel/mL)') + 1
        self.assertEqual([datos.cell(row=k, column=celulas).value for k in (2, 3)], [0.3, 0.5])
        resumen = [fila for fila in libro['Resumen'].iter_rows(values_only=True)]
        self.assertTrue(any(f[0] == 'ABC células' and f[3] == 9.6 for f in resumen))
        self.assertEqual(RegistroAuditoria.objects.filter(categoria='DATOS_PROCESO', accion='GENERAR_INFORME', lote=self.lote).count(), 2)

    def test_la_huella_cambia_con_una_correccion_y_esta_aparece_firmada(self):
        from .services import datos_informe
        antes = datos_informe(self.lote)
        self.assertEqual(antes['huella'], datos_informe(self.lote)['huella'])  # mismos datos, misma huella
        m24 = self.muestreo(4)
        toma = TomaDatos.objects.get(fase=m24)
        r = self.admin.post(f'/api/captura/{m24.id}/corregir/', {'toma': toma.id, 'parametro': self.celulas.id, 'valor': '0.55',
                                                                 'motivo': 'Recuento repetido en cámara', 'firma_password': 'clave-de-prueba'},
                            format='json')
        self.assertEqual(r.status_code, 200, r.data)
        despues = datos_informe(self.lote)
        self.assertNotEqual(antes['huella'], despues['huella'])
        correccion = despues['correcciones'][0]
        self.assertEqual(correccion['cambios'], [{'codigo': 'celulas', 'antes': '0.5', 'despues': '0.55'}])
        self.assertEqual(correccion['firma']['significado'], 'Corrijo el dato de proceso registrado')
        fila = next(t for t in despues['etapas'][0]['tomas'] if t['toma'] == toma.id)
        self.assertTrue(fila['valores'][self.celulas.id]['corregido'])

    def test_informe_por_etapa_con_graficas_combinadas(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from .services import datos_informe
        clave = lambda p: f"{self.etapa.id}-{p.id}"
        datos = datos_informe(self.lote, combinar=[clave(self.celulas), clave(self.temp), clave(self.aspecto), 'basura'],
                              vista='superpuestas')
        self.assertEqual([e['etapa'] for e in datos['etapas']], ['PO4'])
        po4 = datos['etapas'][0]
        self.assertEqual(len(po4['tomas']), 2)
        self.assertEqual([b['titulo'] for b in po4['bloques']], ['[T2] Incubación · Reactor', '[T2] Incubación · Metabolitos',
                                                                 '[T2] Incubación · Cuenta'])
        # Solo numéricos con datos; lo demás se ignora
        self.assertEqual([c['columna']['codigo'] for c in datos['combinadas']], ['celulas', 'temp'])
        self.assertEqual(datos['vista_combinadas'], 'superpuestas')
        r = self.web.get(f'/api/informes-proceso/{self.lote.id}/pdf/?combinar={clave(self.celulas)},{clave(self.temp)}&vista=apiladas')
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b'%PDF'))
        r = self.web.get(f'/api/informes-proceso/{self.lote.id}/excel/?combinar={clave(self.celulas)},{clave(self.temp)}')
        libro = load_workbook(BytesIO(r.content))
        self.assertEqual(libro.sheetnames, ['Resumen', 'Datos PO4', 'Gráficas combinadas', 'Fuera de límites', 'Muestras adicionales',
                                            'Correcciones'])
        self.assertEqual(len(libro['Gráficas combinadas']._charts), 2)
        registro = RegistroAuditoria.objects.filter(accion='GENERAR_INFORME').latest('id')
        self.assertEqual(registro.cambios['combinadas'][1], 'celulas, temp')

    def test_excel_de_varios_lotes_y_validaciones(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.crear_indicadores()
        r = self.web.get(f'/api/informes-proceso/lotes/?proceso={self.proceso.id}&lotes={self.lote.id}')
        self.assertEqual(r.status_code, 200)
        libro = load_workbook(BytesIO(r.content))
        self.assertEqual(libro.sheetnames, ['Información', 'Datos', 'Indicadores'])
        filas = list(libro['Datos'].iter_rows(min_row=2, values_only=True))
        self.assertTrue(any(f[9] == 'celulas' and f[12] == 0.5 for f in filas))
        self.assertEqual(self.web.get(f'/api/informes-proceso/{self.lote.id}/pdf/?etapa=99999').status_code, 404)
        self.assertEqual(self.web.get(f'/api/informes-proceso/lotes/?proceso={self.proceso.id}&lotes=').status_code, 400)


class ImportacionCultivoTests(TestCase):
    """Importación del histórico de Cultivo con datos en memoria (sin PostgreSQL)."""

    def setUp(self):
        from . import importacion_cultivo as imp
        self.imp = imp
        self.t0 = timezone.make_aware(timezone.datetime(2026, 5, 1, 8, 0))
        self.pr1 = self.t0 + timedelta(hours=24)

    def registro(self, id_, etapa, edad, P=True, GyL=False, C=False, **valores):
        inicio = self.t0 if etapa == 'PO4' else self.pr1
        base = {'id': id_, 'lote_id': 'EMAS99001', 'etapa': etapa, 'edad': edad, 'muestra': f'Edad: {edad}',
                'fecha_obj': inicio + timedelta(hours=edad), 'registro': 'E.Prueba',
                'estatus_P': 'registrado' if P else 'Por_registar', 'estatus_GyL': 'registrado' if GyL else 'Por_registar',
                'estatus_C': 'registrado' if C else 'Por_registar'}
        base.update({col: None for col in self.imp.PARAMETROS})
        base.update({k: Decimal(str(v)) if isinstance(v, (int, float)) else v for k, v in valores.items()})
        return base

    def datos(self, status='Finalizado', termino_h=96, extra=()):
        registros = [
            self.registro(1, 'PO4', 0, GyL=True, C=True, temperatura=36.5, glucosa=4.0, celulas=0.5, viabilidad=98,
                          muestra='PO41-M01', lote_medio_de_cultivo='MC-1'),
            self.registro(2, 'PO4', 6, temperatura=36.6, state3='Se ajusta el pH'),
            self.registro(3, 'PO4', 12, P=False, GyL=True, glucosa=3.5),
            self.registro(4, 'PO4', 18, P=False),                                   # sin capturar: no se importa
            self.registro(5, 'PR1', 0, C=True, temperatura=36.8, celulas=1.5),
            self.registro(6, 'PR1', 0, C=True, temperatura=36.8, celulas=1.6),      # duplicado de Cultivo
            self.registro(7, 'PR1', 24, C=True, temperatura=36.9, celulas=3.5),
            *extra,
        ]
        lote = {'lote': 'EMAS99001', 'status': status, 'fecha_inicio': self.t0, 'fecha_inicio_PR1': self.pr1,
                'fecha_termino': self.t0 + timedelta(hours=termino_h), 'tipoReactor': 'Applikon'}
        return {'lotes': [lote], 'registros': {'EMAS99001': registros}}

    def test_crea_receta_lote_y_tomas_con_las_edades_de_cultivo(self):
        resumen = self.imp.importar(self.datos(), curva=[[0, 0.3], [48, 2.0]])
        self.assertTrue(resumen['receta_creada'])
        r = resumen['lotes'][0]
        self.assertEqual((r['estado'], r['tomas']), ('COMPLETADO', 5))
        self.assertEqual(len(r['diferencias']), 1)   # el duplicado con otra cuenta celular se reporta
        lote = LoteProduccion.objects.get(identificador_lote='EMAS99001')
        self.assertEqual(lote.proceso_maestro.nombre, self.imp.NOMBRE_RECETA)
        from .services import tabla_lote
        tabla = tabla_lote(lote)
        edades = [(f['etapa'], f['edad_etapa_h'], f['edad_lote_h']) for f in tabla['filas']]
        self.assertEqual(edades, [('PO4', 0, 0), ('PO4', 6, 6), ('PO4', 12, 12), ('PR1', 0, 24), ('PR1', 24, 48)])
        primera = tabla['filas'][0]
        self.assertEqual(primera['valores']['celulas']['valor'], '0.5')
        self.assertEqual(primera['valores']['lote_medio']['valor'], 'MC-1')
        self.assertEqual(primera['codigos_muestra'], {'Glucosa y lactato': 'PO41-M01'})
        self.assertEqual(tabla['filas'][1]['valores']['observaciones']['valor'], 'Se ajusta el pH')
        self.assertEqual(tabla['filas'][3]['valores']['celulas']['valor'], '1.6')   # gana el registro más reciente
        toma = TomaDatos.objects.filter(fase__lote=lote).first()
        self.assertEqual(toma.registrado_por_nombre, 'E.Prueba (Cultivo)')
        self.assertTrue(RegistroAuditoria.objects.filter(accion='IMPORTAR_HISTORICO', lote=lote).exists())

    def test_indicadores_del_carnet(self):
        from .services import calcular_indicadores
        resumen = self.imp.importar(self.datos(), curva=[[0, 0.3], [48, 2.0]])
        lote = LoteProduccion.objects.get(identificador_lote='EMAS99001')
        ind = calcular_indicadores(lote, resumen['receta'].indicadores.all())
        # ABC por trapecios contra la edad global: (0.5+1.6)/2·24 + (1.6+3.5)/2·24 = 25.2 + 61.2
        self.assertEqual(ind['abc_cel']['valor'], 86.4)
        self.assertEqual(ind['gramos_obj']['valor'], round(86.4 / 1000 - 5.5, 1))
        self.assertIsNotNone(ind['apego_cel']['valor'])

    def test_reimportar_solo_agrega_lo_nuevo_y_actualiza_el_estado(self):
        en_curso = self.datos(status='PR1', termino_h=2000)
        self.imp.importar(en_curso)
        lote = LoteProduccion.objects.get(identificador_lote='EMAS99001')
        self.assertEqual(lote.estado, 'EN_PROGRESO')
        pr1 = lote.fases.get(es_subtarea_muestreo=False, operacion_maestra__etapa__nombre='PR1')
        self.assertEqual(pr1.estado, 'EN_PROGRESO')
        finalizado = self.datos(extra=[self.registro(8, 'PR1', 48, C=True, temperatura=37.0, celulas=5.0)])
        resumen = self.imp.importar(finalizado)
        r = resumen['lotes'][0]
        self.assertFalse(resumen['receta_creada'])
        self.assertEqual((r['nuevo'], r['tomas']), (False, 1))
        lote.refresh_from_db()
        self.assertEqual(lote.estado, 'COMPLETADO')
        self.assertEqual(TomaDatos.objects.filter(fase__lote=lote).count(), 6)
        pr1.refresh_from_db()
        self.assertEqual(pr1.estado, 'COMPLETADA')

    def test_no_toca_un_lote_del_mes_con_el_mismo_codigo(self):
        otra = ProcesoMaestro.objects.create(nombre='Otra receta')
        LoteProduccion.objects.create(identificador_lote='EMAS99001', proceso_maestro=otra, fecha_inicio_planeada=self.t0)
        r = self.imp.importar(self.datos())['lotes'][0]
        self.assertIn('Ya existe', r['aviso'])
        self.assertFalse(TomaDatos.objects.exists())
