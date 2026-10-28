# eventos/tests.py — Eventos GMP: reporte, automáticos, evaluación con firma, archivo del lote e informe.

from datetime import timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from auditoria.models import RegistroAuditoria
from datos_proceso.informes import excel_lote, pdf_lote
from datos_proceso.models import TomaDatos
from datos_proceso.services import datos_informe
from datos_proceso.tests import CapturaBase, usuario_con_rol

from .models import EventoGMP

CLAVE = 'clave-de-prueba'


class EventosBase(CapturaBase):

    def setUp(self):
        super().setUp()
        self.supervisora = usuario_con_rol('supervisora', 'Jefatura')
        self.web = APIClient()
        self.web.force_authenticate(self.supervisora)
        self.calidad = APIClient()
        self.calidad.force_authenticate(usuario_con_rol('calidad', 'Calidad'))
        self.consulta = APIClient()
        self.consulta.force_authenticate(usuario_con_rol('consulta', 'Consulta'))
        self.hito = timezone.now() - timedelta(hours=12, minutes=30)
        self.forzar(self.actividad(self.t1), 'COMPLETADA', self.hito)
        self.forzar(self.actividad(self.t2), 'EN_PROGRESO', self.hito)
        self.m12 = self.forzar(self.muestreo(2), 'EN_PROGRESO')

    def capturar(self, temp):
        r = self.registrar(self.web, self.m12, {**self.valores_reactor(temp=temp), self.glucosa.id: '4'})
        self.assertEqual(r.status_code, 201, r.data)
        return r

    def reportar(self, cliente, **datos):
        cuerpo = {'lote': self.lote.id, 'tipo': 'FALLA_EQUIPO', 'titulo': 'Falla del agitador',
                  'descripcion': 'El agitador se detuvo 10 minutos por una alarma de sobrecorriente.', **datos}
        return cliente.post('/api/eventos/', cuerpo, format='json')

    def cerrar(self, cliente, evento, **datos):
        cuerpo = {'clasificacion': 'SIN_IMPACTO', 'evaluacion': 'Se revisó la tendencia: sin impacto en el lote.',
                  'firma_password': CLAVE, **datos}
        return cliente.post(f'/api/eventos/{evento.id}/cerrar/', cuerpo, format='json')


class EventosAutomaticosTests(EventosBase):

    def test_dato_fuera_de_especificacion_abre_evento(self):
        r = self.capturar('38.4')
        evento = EventoGMP.objects.get()
        self.assertEqual((evento.tipo, evento.origen, evento.estado), ('FUERA_ESPECIFICACION', 'AUTOMATICO', 'ABIERTO'))
        self.assertEqual(evento.valor.parametro, self.temp)
        self.assertEqual(evento.fase, self.m12)
        self.assertEqual(evento.reportado_por_usuario, self.supervisora)   # quien capturó el dato
        self.assertRegex(evento.folio, r'^EV-\d{4}-0001$')
        self.assertIn('36 a 38', evento.descripcion)
        self.assertTrue(any(evento.folio in a for a in r.data['avisos']))
        self.assertTrue(RegistroAuditoria.objects.filter(categoria='EVENTO', accion='EVENTO_AUTOMATICO', lote=self.lote).exists())

    def test_dentro_de_especificacion_no_abre_evento(self):
        self.capturar('37.0')
        self.assertFalse(EventoGMP.objects.exists())

    def test_correccion_fuera_abre_uno_y_no_duplica(self):
        self.capturar('37.0')
        toma = TomaDatos.objects.get(fase=self.m12)

        def corregir(valor):
            return self.web.post(f'/api/captura/{self.m12.id}/corregir/', {
                'toma': toma.id, 'parametro': self.temp.id, 'valor': valor, 'motivo': 'Error de transcripción del termómetro',
                'firma_password': CLAVE}, format='json')

        r = corregir('39')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(EventoGMP.objects.count(), 1)
        self.assertIn('Valor corregido', EventoGMP.objects.get().descripcion)
        self.assertEqual(corregir('39.5').status_code, 200)
        self.assertEqual(EventoGMP.objects.count(), 1)   # sigue abierto el del mismo dato

    def test_inicio_fuera_de_ventana_abre_evento(self):
        t3 = self.actividad(self.t3)
        t2 = self.actividad(self.t2)
        self.forzar(t2, 'COMPLETADA', timezone.now() - timedelta(hours=6))
        t3.inicio_programado = timezone.now() - timedelta(hours=5)
        t3.fin_programado = t3.inicio_programado + timedelta(hours=1)
        t3.save()
        r = self.web.post(f'/api/fases-lote/{t3.id}/cambiar-estado/', {
            'estado': 'EN_PROGRESO', 'motivo_desviacion': 'Se esperó la liberación del cuarto limpio', 'firma_password': CLAVE},
            format='json')
        self.assertEqual(r.status_code, 200, r.data)
        evento = EventoGMP.objects.get(tipo='INICIO_FUERA_VENTANA')
        self.assertEqual(r.data['evento_gmp'], evento.folio)
        self.assertEqual(evento.fase, t3)
        self.assertIn('cuarto limpio', evento.descripcion)


class ReporteYEvaluacionTests(EventosBase):

    def test_reporta_cualquiera_con_sesion_y_valida(self):
        r = self.reportar(self.consulta, fase=self.m12.id, lote='')
        self.assertEqual(r.status_code, 201, r.data)
        evento = EventoGMP.objects.get()
        self.assertEqual((evento.lote, evento.fase, evento.origen), (self.lote, self.m12, 'MANUAL'))
        self.assertTrue(RegistroAuditoria.objects.filter(categoria='EVENTO', accion='REPORTAR_EVENTO').exists())

        self.assertEqual(self.reportar(self.consulta, descripcion='corto').status_code, 400)
        self.assertEqual(self.reportar(self.consulta, tipo='NO_EXISTE').status_code, 400)
        futuro = (timezone.now() + timedelta(hours=2)).isoformat()
        self.assertIn('futura', self.reportar(self.consulta, ocurrio_en=futuro).data['error'])
        self.assertIn(APIClient().post('/api/eventos/', {}, format='json').status_code, (401, 403))   # sin sesión

    def test_reporta_el_operador_desde_la_terminal(self):
        movil = APIClient()
        movil.post('/api/mobile/login/', {'codigo_empleado': 'OP-1', 'pin_acceso': '1234'}, format='json')
        r = self.reportar(movil, fase=self.m12.id, lote='', tipo='PROCEDIMIENTO', titulo='Muestra sin rotular')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(EventoGMP.objects.get().reportado_por_operador, self.operador)

    def test_cierre_con_permiso_firma_y_cuatro_ojos(self):
        evento = EventoGMP.objects.get(pk=self.reportar(self.web).data['id'])
        self.assertEqual(self.cerrar(self.consulta, evento).status_code, 403)                 # sin permiso
        self.assertIn('cuatro ojos', self.cerrar(self.web, evento).data['error'])              # quien lo reportó
        self.assertIn('número', self.cerrar(self.calidad, evento, clasificacion='MAYOR').data['error'])
        r = self.cerrar(self.calidad, evento, firma_password='otra')
        self.assertEqual(r.data['codigo'], 'FIRMA_INVALIDA')
        r = self.cerrar(self.calidad, evento, clasificacion='MENOR', numero_desviacion='DV-2026-031')
        self.assertEqual(r.status_code, 200, r.data)
        evento.refresh_from_db()
        self.assertEqual((evento.estado, evento.clasificacion, evento.numero_desviacion), ('CERRADO', 'MENOR', 'DV-2026-031'))
        self.assertEqual(r.data['firma']['significado'], 'Evalúo y cierro el evento GMP')
        self.assertIn('ya está cerrado', self.cerrar(self.calidad, evento).data['error'])

    def test_anular_con_motivo_y_firma(self):
        evento = EventoGMP.objects.get(pk=self.reportar(self.web).data['id'])
        url = f'/api/eventos/{evento.id}/anular/'
        self.assertEqual(self.calidad.post(url, {'motivo': 'x', 'firma_password': CLAVE}, format='json').status_code, 400)
        r = self.calidad.post(url, {'motivo': 'Se registró en el lote equivocado.', 'firma_password': CLAVE}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        evento.refresh_from_db()
        self.assertEqual(evento.estado, 'ANULADO')
        self.assertTrue(RegistroAuditoria.objects.filter(accion='ANULAR_EVENTO', firma__isnull=False).exists())

    def test_filtros_y_resumen(self):
        self.reportar(self.web)
        self.capturar('38.4')
        self.assertEqual(self.consulta.get('/api/eventos/?estado=ABIERTO').data['total'], 2)
        self.assertEqual(self.consulta.get('/api/eventos/?tipo=FALLA_EQUIPO').data['total'], 1)
        self.assertEqual(self.consulta.get(f'/api/eventos/resumen/?lote={self.lote.id}').data['ABIERTO'], 2)
        self.assertTrue(self.calidad.get('/api/eventos/opciones/').data['puede_evaluar'])
        self.assertFalse(self.consulta.get('/api/eventos/opciones/').data['puede_evaluar'])


class LoteEInformeTests(EventosBase):

    def test_no_se_archiva_con_eventos_abiertos(self):
        evento = EventoGMP.objects.get(pk=self.reportar(self.web).data['id'])
        self.lote.estado = 'COMPLETADO'
        self.lote.save()
        r = self.admin.post(f'/api/lotes/{self.lote.id}/archivar/', {'firma_password': CLAVE}, format='json')
        self.assertEqual(r.data['codigo'], 'EVENTOS_ABIERTOS')
        self.cerrar(self.calidad, evento)
        r = self.admin.post(f'/api/lotes/{self.lote.id}/archivar/', {'firma_password': CLAVE}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertIn('archivado', self.reportar(self.web).data['error'])

    def test_informe_incluye_eventos_y_cambia_la_huella(self):
        evento = EventoGMP.objects.get(pk=self.reportar(self.web, descripcion='Alarma <sobrecorriente> & paro de 10 min').data['id'])
        datos = datos_informe(self.lote)
        self.assertEqual([e['folio'] for e in datos['eventos']], [evento.folio])
        self.assertEqual(datos['eventos'][0]['estado'], 'Abierto')
        self.assertTrue(pdf_lote(datos, 'Prueba').startswith(b'%PDF'))   # el texto con < y & no rompe el PDF
        hoja = load_workbook(BytesIO(excel_lote(datos, 'Prueba')))['Eventos GMP']
        self.assertEqual(hoja.cell(row=2, column=1).value, evento.folio)

        self.cerrar(self.calidad, evento)
        despues = datos_informe(self.lote)
        self.assertEqual(despues['eventos'][0]['firma']['significado'], 'Evalúo y cierro el evento GMP')
        self.assertNotEqual(datos['huella'], despues['huella'])
